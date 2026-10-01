from user.admin.adminSite import adminSite
from django.contrib.auth.decorators import login_required
from django.shortcuts import render
from django.http import HttpResponse
from django.db.models import Sum, Count

from sale.models.sales_receipt import SalesReceipt
from storage.models.silo_movement import SiloMovement
from storage.models.storage_movement import StorageMovement
from parameter.models.campaign import Campaing
from parameter.models.specie import Specie
from user.models.user_producer import UserProducer

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def _producer_filter(user):
    if user.is_superuser or getattr(user, 'type', None) == 'Administration':
        return None
    if getattr(user, 'type', None) == 'Organization':
        return UserProducer.objects.filter(organization=user)
    if getattr(user, 'type', None) == 'Producer':
        return UserProducer.objects.filter(email=user.email)
    if getattr(user, 'type', None) == 'Staff':
        from producer.models.staff import Staff
        staff = Staff.objects.filter(user=user).first()
        if staff:
            return UserProducer.objects.filter(pk=staff.producer_id)
    return UserProducer.objects.none()


ST_OK      = 'ok'
ST_DESVIO  = 'desvio'
ST_ORPHAN  = 'orphan'
ST_PENDING = 'pending'
ST_CANCEL  = 'cancel'

STATUS_LABEL = {
    ST_OK:     'OK',
    ST_DESVIO: 'Desvío',
    ST_ORPHAN: 'Sin salida',
    ST_PENDING:'Pendiente',
    ST_CANCEL: 'Cancelada',
}


def _build_traceability_data(request):
    user = request.user
    prod_filter = _producer_filter(user)
    show_producer = user.is_superuser or getattr(user, 'type', None) == 'Administration'

    campaing_id  = request.GET.get('campaing_id')
    specie_id    = request.GET.get('specie_id')
    state_filter = request.GET.get('state_filter', '')
    date_from    = request.GET.get('date_from')
    date_to      = request.GET.get('date_to')

    # ── Receipts ─────────────────────────────────────────────────
    receipts_qs = SalesReceipt.objects.select_related(
        'type', 'campaing', 'especie', 'currency', 'producer',
    ).order_by('-date', '-id')

    if prod_filter is not None:
        receipts_qs = receipts_qs.filter(producer__in=prod_filter)
    if campaing_id:
        receipts_qs = receipts_qs.filter(campaing_id=campaing_id)
    if specie_id:
        receipts_qs = receipts_qs.filter(especie_id=specie_id)
    if state_filter in ('confirm', 'issues'):
        receipts_qs = receipts_qs.filter(state='confirm')
    if date_from:
        receipts_qs = receipts_qs.filter(date__gte=date_from)
    if date_to:
        receipts_qs = receipts_qs.filter(date__lte=date_to)

    receipts   = list(receipts_qs)
    receipt_ids = [r.pk for r in receipts]

    # Two aggregation queries — no N+1
    silo_by = {
        row['related_sale_id']: {'tons': row['total'] or 0.0, 'count': row['cnt']}
        for row in (
            SiloMovement.objects
            .filter(related_sale_id__in=receipt_ids, type=SiloMovement.EXIT)
            .values('related_sale_id')
            .annotate(total=Sum('quantity_tons'), cnt=Count('id'))
        )
    }
    stor_by = {
        row['related_sale_id']: {'tons': row['total'] or 0.0, 'count': row['cnt']}
        for row in (
            StorageMovement.objects
            .filter(related_sale_id__in=receipt_ids, type=StorageMovement.SALE_EXIT)
            .values('related_sale_id')
            .annotate(total=Sum('quantity_tons'), cnt=Count('id'))
        )
    }

    rows = []
    for rec in receipts:
        sd = silo_by.get(rec.pk, {'tons': 0.0, 'count': 0})
        xd = stor_by.get(rec.pk, {'tons': 0.0, 'count': 0})
        linked = round(sd['tons'] + xd['tons'], 3)
        n_mvmt = sd['count'] + xd['count']

        if rec.state == 'confirm':
            if n_mvmt == 0:
                status = ST_ORPHAN
            elif abs(linked - rec.amount) <= 0.011:
                status = ST_OK
            else:
                status = ST_DESVIO
        elif rec.state == 'cancel':
            status = ST_CANCEL
        else:
            status = ST_PENDING

        if state_filter == 'issues' and status not in (ST_ORPHAN, ST_DESVIO):
            continue

        rows.append({
            'date': rec.date,
            'type': str(rec.type),
            'number': rec.number,
            'buyer': rec.name_custom,
            'campaing': str(rec.campaing),
            'specie': str(rec.especie) if rec.especie else '—',
            'currency': str(rec.currency),
            'amount': rec.amount,
            'silo_tons': round(sd['tons'], 3),
            'stor_tons': round(xd['tons'], 3),
            'linked_tons': linked,
            'mvmt_count': n_mvmt,
            'desvio': round(linked - rec.amount, 3) if rec.state == 'confirm' else None,
            'status': status,
            'status_label': STATUS_LABEL[status],
            'state': rec.state,
            'producer': str(rec.producer) if show_producer else '',
        })

    # ── Orphan exits — movements not linked to any sale ──────────
    orphan_silo_qs = SiloMovement.objects.select_related(
        'silo', 'crop', 'crop__campaing', 'crop__variety__specie', 'producer',
    ).filter(type=SiloMovement.EXIT, related_sale__isnull=True)

    orphan_stor_qs = StorageMovement.objects.select_related(
        'certificate', 'certificate__campaing', 'certificate__specie',
        'certificate__storage_provider', 'producer',
    ).filter(type=StorageMovement.SALE_EXIT, related_sale__isnull=True)

    if prod_filter is not None:
        orphan_silo_qs = orphan_silo_qs.filter(producer__in=prod_filter)
        orphan_stor_qs = orphan_stor_qs.filter(producer__in=prod_filter)
    if campaing_id:
        orphan_silo_qs = orphan_silo_qs.filter(crop__campaing_id=campaing_id)
        orphan_stor_qs = orphan_stor_qs.filter(certificate__campaing_id=campaing_id)
    if specie_id:
        orphan_silo_qs = orphan_silo_qs.filter(crop__variety__specie_id=specie_id)
        orphan_stor_qs = orphan_stor_qs.filter(certificate__specie_id=specie_id)
    if date_from:
        orphan_silo_qs = orphan_silo_qs.filter(date__gte=date_from)
        orphan_stor_qs = orphan_stor_qs.filter(date__gte=date_from)
    if date_to:
        orphan_silo_qs = orphan_silo_qs.filter(date__lte=date_to)
        orphan_stor_qs = orphan_stor_qs.filter(date__lte=date_to)

    orphan_rows = []
    for m in orphan_silo_qs.order_by('-date'):
        try:
            specie_str   = str(m.crop.variety.specie) if (m.crop and m.crop.variety) else '—'
            campaing_str = str(m.crop.campaing) if m.crop else '—'
        except Exception:
            specie_str = campaing_str = '—'
        orphan_rows.append({
            'source': 'silo',
            'source_label': 'Silo propio',
            'date': m.date,
            'location': str(m.silo),
            'campaing': campaing_str,
            'specie': specie_str,
            'quantity_tons': m.quantity_tons,
            'remito': m.remito_number or '',
            'producer': str(m.producer) if show_producer else '',
        })
    for m in orphan_stor_qs.order_by('-date'):
        cert = m.certificate
        orphan_rows.append({
            'source': 'acopio',
            'source_label': 'Acopio tercero',
            'date': m.date,
            'location': f'{cert.number} ({cert.storage_provider})',
            'campaing': str(cert.campaing),
            'specie': str(cert.specie),
            'quantity_tons': m.quantity_tons,
            'remito': m.remito_number or '',
            'producer': str(m.producer) if show_producer else '',
        })
    orphan_rows.sort(key=lambda r: r['date'], reverse=True)

    counts = {st: 0 for st in (ST_OK, ST_DESVIO, ST_ORPHAN, ST_PENDING, ST_CANCEL)}
    for r in rows:
        counts[r['status']] += 1

    return {
        'rows': rows,
        'orphan_rows': orphan_rows,
        'counts': counts,
        'show_producer': show_producer,
        'campaigns': Campaing.objects.order_by('-name'),
        'species': Specie.objects.order_by('name'),
        'selected_campaing_id':  campaing_id   or '',
        'selected_specie_id':    specie_id     or '',
        'selected_state_filter': state_filter  or '',
        'selected_date_from':    date_from     or '',
        'selected_date_to':      date_to       or '',
    }


@login_required
def traceability_view(request):
    ctx = _build_traceability_data(request)
    return render(request, 'sale/report_traceability.html', {**adminSite.each_context(request), **ctx})


@login_required
def traceability_excel(request):
    ctx = _build_traceability_data(request)

    wb = openpyxl.Workbook()
    hdr_font  = Font(bold=True, color='FFFFFF', size=10)
    hdr_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    num_fmt = '#,##0.000'

    # ── Sheet 1: Ventas ──────────────────────────────────────────
    ws1 = wb.active
    ws1.title = 'Ventas'
    grn_fill = PatternFill('solid', fgColor='1E7D3B')

    headers1 = [
        'Fecha', 'Tipo', 'Número', 'Comprador', 'Campaña', 'Especie',
        'Moneda', 'Tn venta', 'Tn silo', 'Tn acopio',
        'Tn total salida', 'Desvío', 'Movimientos', 'Estado',
    ]
    if ctx['show_producer']:
        headers1.append('Productor')

    for col_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col_idx, value=h)
        cell.fill = grn_fill; cell.font = hdr_font; cell.alignment = hdr_align

    for ri, r in enumerate(ctx['rows'], 2):
        ws1.cell(ri, 1,  str(r['date']))
        ws1.cell(ri, 2,  r['type'])
        ws1.cell(ri, 3,  r['number'])
        ws1.cell(ri, 4,  r['buyer'])
        ws1.cell(ri, 5,  r['campaing'])
        ws1.cell(ri, 6,  r['specie'])
        ws1.cell(ri, 7,  r['currency'])
        ws1.cell(ri, 8,  r['amount']).number_format = num_fmt
        ws1.cell(ri, 9,  r['silo_tons']).number_format = num_fmt
        ws1.cell(ri, 10, r['stor_tons']).number_format = num_fmt
        ws1.cell(ri, 11, r['linked_tons']).number_format = num_fmt
        ws1.cell(ri, 12, r['desvio']).number_format = num_fmt
        ws1.cell(ri, 13, r['mvmt_count'])
        ws1.cell(ri, 14, r['status_label'])
        if ctx['show_producer']:
            ws1.cell(ri, 15, r['producer'])

    ws1.freeze_panes = 'A2'
    for col in ws1.columns:
        w = max((len(str(c.value or '')) for c in col), default=8)
        ws1.column_dimensions[get_column_letter(col[0].column)].width = min(w + 2, 38)

    # ── Sheet 2: Salidas sin venta ───────────────────────────────
    ws2 = wb.create_sheet('Salidas sin venta')
    red_fill = PatternFill('solid', fgColor='c0392b')

    headers2 = ['Fecha', 'Origen', 'Ubicación', 'Campaña', 'Especie', 'Tn', 'Remito']
    if ctx['show_producer']:
        headers2.append('Productor')

    for col_idx, h in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col_idx, value=h)
        cell.fill = red_fill; cell.font = hdr_font; cell.alignment = hdr_align

    for ri, r in enumerate(ctx['orphan_rows'], 2):
        ws2.cell(ri, 1, str(r['date']))
        ws2.cell(ri, 2, r['source_label'])
        ws2.cell(ri, 3, r['location'])
        ws2.cell(ri, 4, r['campaing'])
        ws2.cell(ri, 5, r['specie'])
        ws2.cell(ri, 6, r['quantity_tons']).number_format = num_fmt
        ws2.cell(ri, 7, r['remito'])
        if ctx['show_producer']:
            ws2.cell(ri, 8, r['producer'])

    ws2.freeze_panes = 'A2'
    for col in ws2.columns:
        w = max((len(str(c.value or '')) for c in col), default=8)
        ws2.column_dimensions[get_column_letter(col[0].column)].width = min(w + 2, 38)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="trazabilidad_ventas.xlsx"'
    wb.save(response)
    return response
