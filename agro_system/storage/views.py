from user.admin.adminSite import adminSite
from datetime import date

from django.contrib.auth.decorators import login_required
from django.shortcuts import render
from django.http import HttpResponse
from collections import defaultdict

from django.db.models import Sum, Q, Count
from django.db.models.functions import Coalesce

from storage.models.silo_movement import SiloMovement
from storage.models.storage_movement import StorageMovement
from storage.models.silo import Silo
from storage.models.silo_stock import SiloStock
from storage.models.storage_certificate import StorageCertificate
from storage.models.storage_service import StorageService
from parameter.models.campaign import Campaing
from parameter.models.specie import Specie
from parameter.models.storage_service_type import StorageServiceType
from parameter.models.currency import Currency
from operation.models.crop import Crop
from operation.models.crop_yields import CropYields
from input.models.provider import Provider
from user.models.user_producer import UserProducer

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def _producer_filter(user):
    if user.is_superuser or getattr(user, 'type', None) == 'Administration':
        return None
    if getattr(user, 'type', None) == 'Organization':
        return UserProducer.objects.filter(organization=user)
    if getattr(user, 'type', None) == 'Producer':
        return UserProducer.objects.filter(email=user.email)
    if getattr(user, 'type', None) == 'Staff':
        from producer.models.staff import Staff
        staff = Staff.objects.filter(user=user).first()
        if staff:
            return UserProducer.objects.filter(pk=staff.producer_id)
    return UserProducer.objects.none()


def _build_acopio_data(request):
    user = request.user
    prod_filter = _producer_filter(user)
    show_producer = user.is_superuser or getattr(user, 'type', None) == 'Administration'

    campaing_id = request.GET.get('campaing_id')
    crop_id = request.GET.get('crop_id')
    silo_id = request.GET.get('silo_id')
    certificate_id = request.GET.get('certificate_id')
    source = request.GET.get('source', '')
    type_filter = request.GET.get('type_filter', '')
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')

    def _apply_prod(qs, field='producer'):
        if prod_filter is not None:
            qs = qs.filter(**{f'{field}__in': prod_filter})
        return qs

    # ── Silo movements ──────────────────────────────────────────────
    silo_rows = []
    if source in ('', 'silo'):
        sm_qs = SiloMovement.objects.select_related(
            'silo', 'silo__countryside',
            'crop', 'crop__campaing', 'crop__activity', 'crop__activity__specie',
            'destination_silo', 'driver', 'related_sale',
        ).order_by('-date', '-id')
        sm_qs = _apply_prod(sm_qs, 'producer')
        if campaing_id:
            sm_qs = sm_qs.filter(crop__campaing_id=campaing_id)
        if crop_id:
            sm_qs = sm_qs.filter(crop_id=crop_id)
        if silo_id:
            sm_qs = sm_qs.filter(silo_id=silo_id)
        if date_from:
            sm_qs = sm_qs.filter(date__gte=date_from)
        if date_to:
            sm_qs = sm_qs.filter(date__lte=date_to)
        if type_filter == 'entry':
            sm_qs = sm_qs.filter(type='entry')
        elif type_filter == 'exit':
            sm_qs = sm_qs.filter(type__in=['exit', 'transfer'])

        for m in sm_qs[:500]:
            dest = ''
            if m.type == 'transfer' and m.destination_silo:
                dest = f'→ {m.destination_silo}'
            try:
                specie_str = str(m.crop.activity.specie) if (m.crop and m.crop.activity) else '—'
            except Exception:
                specie_str = '—'
            silo_rows.append({
                'source': 'silo',
                'source_label': 'Silo propio',
                'date': m.date,
                'type': m.type,
                'type_display': m.get_type_display(),
                'campaing': str(m.crop.campaing) if m.crop else '—',
                'crop': str(m.crop) if m.crop else '—',
                'specie': specie_str,
                'location': str(m.silo),
                'identifier': str(m.silo),
                'quantity_tons': m.quantity_tons,
                'remito': m.remito_number or '',
                'driver': str(m.driver) if m.driver else '',
                'destination': dest,
                'related_sale': str(m.related_sale) if m.related_sale else '',
                'notes': m.notes or '',
                'producer': str(m.producer) if show_producer else '',
            })

    # ── Storage movements ────────────────────────────────────────────
    stor_rows = []
    if source in ('', 'acopio'):
        stm_qs = StorageMovement.objects.select_related(
            'certificate', 'certificate__campaing', 'certificate__specie',
            'certificate__crop', 'certificate__storage_provider',
            'driver', 'related_sale', 'port',
        ).order_by('-date', '-id')
        stm_qs = _apply_prod(stm_qs, 'producer')
        if campaing_id:
            stm_qs = stm_qs.filter(certificate__campaing_id=campaing_id)
        if crop_id:
            stm_qs = stm_qs.filter(certificate__crop_id=crop_id)
        if certificate_id:
            stm_qs = stm_qs.filter(certificate_id=certificate_id)
        if date_from:
            stm_qs = stm_qs.filter(date__gte=date_from)
        if date_to:
            stm_qs = stm_qs.filter(date__lte=date_to)
        if type_filter == 'entry':
            stm_qs = stm_qs.filter(type='entry')
        elif type_filter == 'exit':
            stm_qs = stm_qs.filter(type__in=['sale_exit', 'withdrawal', 'transfer'])

        for m in stm_qs[:500]:
            cert = m.certificate
            dest = m.destination or ''
            if m.port:
                dest = f'{dest} / Puerto: {m.port}' if dest else f'Puerto: {m.port}'
            stor_rows.append({
                'source': 'acopio',
                'source_label': 'Acopio',
                'date': m.date,
                'type': m.type,
                'type_display': m.get_type_display(),
                'campaing': str(cert.campaing),
                'crop': str(cert.crop) if cert.crop else '—',
                'specie': str(cert.specie),
                'location': str(cert.storage_provider),
                'identifier': f'{cert.number} ({cert.storage_provider})',
                'quantity_tons': m.quantity_tons,
                'remito': m.remito_number or '',
                'driver': str(m.driver) if m.driver else '',
                'destination': dest,
                'related_sale': str(m.related_sale) if m.related_sale else '',
                'notes': m.notes or '',
                'producer': str(m.producer) if show_producer else '',
            })

    # Merge and sort by date desc
    all_rows = silo_rows + stor_rows
    all_rows.sort(key=lambda r: r['date'], reverse=True)
    truncated = len(all_rows) > 1000
    all_rows = all_rows[:1000]

    # Summary stats
    entry_types = {'entry'}
    exit_types = {'exit', 'sale_exit', 'withdrawal', 'transfer'}
    tons_entry = sum(r['quantity_tons'] for r in all_rows if r['type'] in entry_types)
    tons_exit = sum(r['quantity_tons'] for r in all_rows if r['type'] in exit_types)
    tons_net = tons_entry - tons_exit

    # Filter option querysets
    campaigns_qs = Campaing.objects.order_by('-name')

    crops_qs = Crop.objects.select_related('campaing', 'activity').order_by('name')
    if campaing_id:
        crops_qs = crops_qs.filter(campaing_id=campaing_id)
    else:
        crops_qs = crops_qs.none()

    silos_qs = _apply_prod(
        Silo.objects.select_related('countryside').order_by('name'), 'producer'
    )

    certs_qs = StorageCertificate.objects.select_related(
        'storage_provider', 'campaing'
    ).order_by('-date_issue')
    certs_qs = _apply_prod(certs_qs, 'producer')
    if campaing_id:
        certs_qs = certs_qs.filter(campaing_id=campaing_id)

    selected_campaing = Campaing.objects.filter(pk=campaing_id).first() if campaing_id else None

    return {
        'rows': all_rows,
        'truncated': truncated,
        'tons_entry': tons_entry,
        'tons_exit': tons_exit,
        'tons_net': tons_net,
        'campaigns': campaigns_qs,
        'all_crops': list(crops_qs),
        'all_silos': list(silos_qs),
        'all_certificates': list(certs_qs[:150]),
        'selected_campaing_id': campaing_id or '',
        'selected_campaing': selected_campaing,
        'selected_crop_id': crop_id or '',
        'selected_silo_id': silo_id or '',
        'selected_certificate_id': certificate_id or '',
        'selected_source': source,
        'selected_type_filter': type_filter,
        'selected_date_from': date_from or '',
        'selected_date_to': date_to or '',
        'show_producer': show_producer,
    }


@login_required
def acopio_ledger_view(request):
    ctx = _build_acopio_data(request)
    return render(request, 'storage/report_ledger.html', {**adminSite.each_context(request), **ctx})


@login_required
def acopio_ledger_excel(request):
    ctx = _build_acopio_data(request)
    rows = ctx['rows']

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Ledger Acopio'

    hdr_fill_green = PatternFill('solid', fgColor='1E7D3B')
    hdr_fill_blue = PatternFill('solid', fgColor='2c6e9e')
    hdr_font = Font(bold=True, color='FFFFFF', size=10)
    hdr_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    headers = [
        'Fecha', 'Origen', 'Silo / Certificado', 'Campaña', 'Cultivo', 'Especie',
        'Tipo movimiento', 'Toneladas', 'Remito', 'Chofer', 'Destino',
        'Venta relacionada', 'Notas',
    ]
    if ctx['show_producer']:
        headers.append('Productor')

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = hdr_fill_green
        cell.font = hdr_font
        cell.alignment = hdr_align

    for row_idx, r in enumerate(rows, 2):
        ws.cell(row=row_idx, column=1, value=r['date'])
        ws.cell(row=row_idx, column=1).number_format = 'DD/MM/YYYY'
        ws.cell(row=row_idx, column=2, value=r['source_label'])
        # Color-code source column
        if r['source'] == 'acopio':
            ws.cell(row=row_idx, column=2).fill = PatternFill('solid', fgColor='dce8f5')
        ws.cell(row=row_idx, column=3, value=r['identifier'])
        ws.cell(row=row_idx, column=4, value=r['campaing'])
        ws.cell(row=row_idx, column=5, value=r['crop'])
        ws.cell(row=row_idx, column=6, value=r['specie'])
        ws.cell(row=row_idx, column=7, value=r['type_display'])
        ws.cell(row=row_idx, column=8, value=r['quantity_tons'])
        ws.cell(row=row_idx, column=8).number_format = '#,##0.000'
        ws.cell(row=row_idx, column=9, value=r['remito'])
        ws.cell(row=row_idx, column=10, value=r['driver'])
        ws.cell(row=row_idx, column=11, value=r['destination'])
        ws.cell(row=row_idx, column=12, value=r['related_sale'])
        ws.cell(row=row_idx, column=13, value=r['notes'])
        if ctx['show_producer']:
            ws.cell(row=row_idx, column=14, value=r['producer'])

    ws.freeze_panes = 'A2'

    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 45)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="ledger_acopio.xlsx"'
    wb.save(response)
    return response


# ─────────────────────────────────────────────────────────────────────────────
# Posición de stock acopiado
# ─────────────────────────────────────────────────────────────────────────────

def _age_class(age_days):
    if age_days is None:
        return ''
    if age_days <= 30:
        return 'age-fresh'
    if age_days <= 90:
        return 'age-moderate'
    if age_days <= 180:
        return 'age-old'
    return 'age-very-old'


def _pct_class(pct):
    if pct is None:
        return ''
    if pct <= 50:
        return 'pct-low'
    if pct <= 85:
        return 'pct-mid'
    return 'pct-high'


def _build_stock_position_data(request):
    user = request.user
    prod_filter = _producer_filter(user)
    show_producer = user.is_superuser or getattr(user, 'type', None) == 'Administration'

    campaing_id = request.GET.get('campaing_id')
    specie_id = request.GET.get('specie_id')
    source = request.GET.get('source', '')
    only_active = request.GET.get('only_active', '1')

    today = date.today()

    def _ap(qs, field='producer'):
        if prod_filter is not None:
            qs = qs.filter(**{f'{field}__in': prod_filter})
        return qs

    # ── Silos propios (via SiloStock) ────────────────────────────────
    silo_groups = []
    if source in ('', 'silo'):
        ss_qs = SiloStock.objects.select_related(
            'silo', 'silo__countryside', 'silo__type',
            'specie', 'campaing', 'crop',
        ).order_by('silo__countryside__name', 'silo__name', '-campaing__name', 'specie__name')
        ss_qs = _ap(ss_qs, 'silo__producer')
        if campaing_id:
            ss_qs = ss_qs.filter(campaing_id=campaing_id)
        if specie_id:
            ss_qs = ss_qs.filter(specie_id=specie_id)
        if only_active == '1':
            ss_qs = ss_qs.filter(quantity_tons__gt=0)

        silo_dict = {}
        for stock in ss_qs:
            sid = stock.silo_id
            if sid not in silo_dict:
                silo_dict[sid] = {
                    'silo': stock.silo,
                    'total_tons': 0.0,
                    'stocks': [],
                    'producer': str(stock.silo.producer) if show_producer else '',
                }
            age_days = (today - stock.date_entry).days if stock.date_entry else None
            silo_dict[sid]['total_tons'] += stock.quantity_tons
            silo_dict[sid]['stocks'].append({
                'specie': str(stock.specie),
                'campaing': str(stock.campaing),
                'crop': str(stock.crop) if stock.crop else '—',
                'quantity_tons': stock.quantity_tons,
                'humidity_pct': stock.humidity_pct,
                'impurities_pct': stock.impurities_pct,
                'protein_pct': stock.protein_pct,
                'test_weight': stock.test_weight,
                'date_entry': stock.date_entry,
                'age_days': age_days,
                'age_class': _age_class(age_days),
            })

        for sg in silo_dict.values():
            cap = sg['silo'].capacity_tons or 0
            sg['pct_capacity'] = round(sg['total_tons'] / cap * 100, 1) if cap else None
            sg['pct_class'] = _pct_class(sg['pct_capacity'])
            silo_groups.append(sg)

        silo_groups.sort(key=lambda sg: (str(sg['silo'].countryside), sg['silo'].name))

    # ── Certificados de acopio (StorageCertificate) ─────────────────
    certs_list = []
    if source in ('', 'acopio'):
        cert_qs = StorageCertificate.objects.select_related(
            'storage_provider', 'campaing', 'specie', 'crop',
        ).annotate(
            tons_in=Coalesce(
                Sum('movements__quantity_tons', filter=Q(movements__type='entry')), 0.0
            ),
            tons_out=Coalesce(
                Sum('movements__quantity_tons',
                    filter=Q(movements__type__in=['sale_exit', 'withdrawal', 'transfer'])),
                0.0,
            ),
        ).order_by('-date_issue')
        cert_qs = _ap(cert_qs, 'producer')
        if campaing_id:
            cert_qs = cert_qs.filter(campaing_id=campaing_id)
        if specie_id:
            cert_qs = cert_qs.filter(specie_id=specie_id)
        if only_active == '1':
            cert_qs = cert_qs.exclude(state=StorageCertificate.CANCELLED)

        for cert in cert_qs:
            qty_avail = max(cert.tons_in - cert.tons_out, 0)
            pct_remaining = (
                round(qty_avail / cert.quantity_entry * 100, 1)
                if cert.quantity_entry else None
            )
            age_days = (today - cert.date_issue).days if cert.date_issue else None
            certs_list.append({
                'cert': cert,
                'qty_available': qty_avail,
                'qty_entry': cert.quantity_entry,
                'pct_remaining': pct_remaining,
                'age_days': age_days,
                'age_class': _age_class(age_days),
                'state': cert.get_state_display(),
                'state_raw': cert.state,
                'is_warrant': cert.is_warrant,
                'producer': str(cert.producer) if show_producer else '',
            })

    total_silo_tons = sum(sg['total_tons'] for sg in silo_groups)
    total_cert_tons = sum(c['qty_available'] for c in certs_list)

    return {
        'silo_groups': silo_groups,
        'certs_list': certs_list,
        'total_silo_tons': total_silo_tons,
        'total_cert_tons': total_cert_tons,
        'total_tons': total_silo_tons + total_cert_tons,
        'campaigns': Campaing.objects.order_by('-name'),
        'species': Specie.objects.order_by('name'),
        'selected_campaing_id': campaing_id or '',
        'selected_specie_id': specie_id or '',
        'selected_source': source,
        'only_active': only_active,
        'show_producer': show_producer,
        'today': today,
    }


@login_required
def stock_position_view(request):
    ctx = _build_stock_position_data(request)
    return render(request, 'storage/report_stock_position.html', {**adminSite.each_context(request), **ctx})


@login_required
def stock_position_excel(request):
    ctx = _build_stock_position_data(request)

    wb = openpyxl.Workbook()

    hdr_font = Font(bold=True, color='FFFFFF', size=10)
    hdr_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    def make_hdr(ws, headers, color):
        fill = PatternFill('solid', fgColor=color)
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = fill
            cell.font = hdr_font
            cell.alignment = hdr_align

    def autowidth(ws):
        for col in ws.columns:
            max_len = max((len(str(c.value or '')) for c in col), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 40)

    # ── Sheet 1: Silos ──
    ws1 = wb.active
    ws1.title = 'Silos propios'
    silo_headers = [
        'Campo', 'Silo', 'Tipo silo', 'Capacidad (t)', 'Especie', 'Campaña', 'Cultivo',
        'Stock (t)', 'Humedad %', 'Impurezas %', 'Proteína %', 'Peso esp.',
        'Último movim.', 'Antigüedad (días)',
    ]
    if ctx['show_producer']:
        silo_headers.append('Productor')
    make_hdr(ws1, silo_headers, '1E7D3B')

    row_idx = 2
    for sg in ctx['silo_groups']:
        silo = sg['silo']
        for st in sg['stocks']:
            ws1.cell(row=row_idx, column=1, value=str(silo.countryside))
            ws1.cell(row=row_idx, column=2, value=silo.name)
            ws1.cell(row=row_idx, column=3, value=str(silo.type))
            ws1.cell(row=row_idx, column=4, value=silo.capacity_tons)
            ws1.cell(row=row_idx, column=4).number_format = '#,##0.000'
            ws1.cell(row=row_idx, column=5, value=st['specie'])
            ws1.cell(row=row_idx, column=6, value=st['campaing'])
            ws1.cell(row=row_idx, column=7, value=st['crop'])
            ws1.cell(row=row_idx, column=8, value=st['quantity_tons'])
            ws1.cell(row=row_idx, column=8).number_format = '#,##0.000'
            ws1.cell(row=row_idx, column=9, value=st['humidity_pct'])
            ws1.cell(row=row_idx, column=10, value=st['impurities_pct'])
            ws1.cell(row=row_idx, column=11, value=st['protein_pct'])
            ws1.cell(row=row_idx, column=12, value=st['test_weight'])
            ws1.cell(row=row_idx, column=13, value=st['date_entry'])
            ws1.cell(row=row_idx, column=13).number_format = 'DD/MM/YYYY'
            ws1.cell(row=row_idx, column=14, value=st['age_days'])
            if ctx['show_producer']:
                ws1.cell(row=row_idx, column=15, value=sg['producer'])
            row_idx += 1
    ws1.freeze_panes = 'A2'
    autowidth(ws1)

    # ── Sheet 2: Certificados ──
    ws2 = wb.create_sheet('Acopio externo')
    cert_headers = [
        'Certificado', 'Acopiador', 'Especie', 'Campaña', 'Cultivo',
        'Ingreso inicial (t)', 'Stock disponible (t)', '% disponible',
        'Estado', 'Warrant', 'Fecha emisión', 'Antigüedad (días)',
    ]
    if ctx['show_producer']:
        cert_headers.append('Productor')
    make_hdr(ws2, cert_headers, '2c6e9e')

    for row_idx, c in enumerate(ctx['certs_list'], 2):
        cert = c['cert']
        ws2.cell(row=row_idx, column=1, value=cert.number)
        ws2.cell(row=row_idx, column=2, value=str(cert.storage_provider))
        ws2.cell(row=row_idx, column=3, value=str(cert.specie))
        ws2.cell(row=row_idx, column=4, value=str(cert.campaing))
        ws2.cell(row=row_idx, column=5, value=str(cert.crop) if cert.crop else '')
        ws2.cell(row=row_idx, column=6, value=c['qty_entry'])
        ws2.cell(row=row_idx, column=6).number_format = '#,##0.000'
        ws2.cell(row=row_idx, column=7, value=c['qty_available'])
        ws2.cell(row=row_idx, column=7).number_format = '#,##0.000'
        ws2.cell(row=row_idx, column=8, value=c['pct_remaining'])
        ws2.cell(row=row_idx, column=9, value=c['state'])
        ws2.cell(row=row_idx, column=10, value='Sí' if c['is_warrant'] else 'No')
        ws2.cell(row=row_idx, column=11, value=cert.date_issue)
        ws2.cell(row=row_idx, column=11).number_format = 'DD/MM/YYYY'
        ws2.cell(row=row_idx, column=12, value=c['age_days'])
        if ctx['show_producer']:
            ws2.cell(row=row_idx, column=13, value=c['producer'])
    ws2.freeze_panes = 'A2'
    autowidth(ws2)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="stock_acopio.xlsx"'
    wb.save(response)
    return response


# ─────────────────────────────────────────────────────────────────────────────
# Conciliación rinde vs. acopiado
# ─────────────────────────────────────────────────────────────────────────────

def _build_conciliation_data(request):
    user = request.user
    prod_filter = _producer_filter(user)
    show_producer = user.is_superuser or getattr(user, 'type', None) == 'Administration'

    campaing_id = request.GET.get('campaing_id')
    specie_id = request.GET.get('specie_id')
    only_with_real = request.GET.get('only_with_real', '')

    # ── Crops ──────────────────────────────────────────────────────────
    crops_qs = Crop.objects.select_related(
        'campaing', 'countryside', 'lot', 'activity', 'activity__specie', 'producer',
    ).order_by('campaing__name', 'countryside__name', 'activity__name')

    if prod_filter is not None:
        crops_qs = crops_qs.filter(producer__in=prod_filter)
    if campaing_id:
        crops_qs = crops_qs.filter(campaing_id=campaing_id)
    if specie_id:
        crops_qs = crops_qs.filter(activity__specie_id=specie_id)

    crops = list(crops_qs)
    crop_ids = [c.pk for c in crops]

    if not crop_ids:
        return _conciliation_empty_ctx(campaing_id, specie_id, only_with_real, show_producer)

    # ── CropYields: most recent estimate and real per crop ─────────────
    yields_map = {}

    # Most recent estimate per crop
    seen = set()
    for y in CropYields.objects.filter(crop_id__in=crop_ids).order_by('crop_id', '-date_estimate', '-pk'):
        if y.crop_id not in seen:
            seen.add(y.crop_id)
            yields_map.setdefault(y.crop_id, {})
            yields_map[y.crop_id]['est'] = y.yields_estimate
            yields_map[y.crop_id]['est_date'] = y.date_estimate

    # Most recent real per crop
    seen = set()
    for y in (CropYields.objects.filter(crop_id__in=crop_ids, yields_real__isnull=False)
              .order_by('crop_id', '-date_real', '-pk')):
        if y.crop_id not in seen:
            seen.add(y.crop_id)
            yields_map.setdefault(y.crop_id, {})
            yields_map[y.crop_id]['real'] = y.yields_real
            yields_map[y.crop_id]['real_date'] = y.date_real

    # ── SiloStock: current stock per crop ──────────────────────────────
    silo_tons = {
        row['crop_id']: row['total']
        for row in (SiloStock.objects
                    .filter(crop_id__in=crop_ids)
                    .values('crop_id')
                    .annotate(total=Sum('quantity_tons')))
    }

    # ── StorageCertificate: available tons per crop ────────────────────
    cert_tons = {}
    for cert in (StorageCertificate.objects
                 .filter(crop_id__in=crop_ids)
                 .annotate(
                     tons_in=Coalesce(Sum('movements__quantity_tons', filter=Q(movements__type='entry')), 0.0),
                     tons_out=Coalesce(Sum('movements__quantity_tons',
                                          filter=Q(movements__type__in=['sale_exit', 'withdrawal', 'transfer'])),
                                       0.0),
                 )):
        avail = max(cert.tons_in - cert.tons_out, 0)
        cert_tons[cert.crop_id] = cert_tons.get(cert.crop_id, 0.0) + avail

    # ── Build rows ─────────────────────────────────────────────────────
    rows = []
    for crop in crops:
        yd = yields_map.get(crop.pk, {})
        ha = crop.hectares_harvested or 0

        est_rate = yd.get('est')
        real_rate = yd.get('real')

        tn_est = round(est_rate * ha, 3) if est_rate is not None else None
        tn_real = round(real_rate * ha, 3) if real_rate is not None else None

        tn_silo = round(silo_tons.get(crop.pk, 0.0), 3)
        tn_cert = round(cert_tons.get(crop.pk, 0.0), 3)
        tn_total = round(tn_silo + tn_cert, 3)

        # Desvío = real - acopiado  (positive → grain unaccounted for)
        desvio = round(tn_real - tn_total, 3) if tn_real is not None else None

        # % conciliado = acopiado / real × 100
        pct = round(tn_total / tn_real * 100, 1) if (tn_real and tn_real > 0) else None

        if only_with_real and real_rate is None:
            continue

        # colour class for pct badge
        if pct is None:
            pct_class = 'pct-unknown'
        elif pct > 100:
            pct_class = 'pct-over'
        elif pct >= 95:
            pct_class = 'pct-ok'
        elif pct >= 80:
            pct_class = 'pct-warn'
        else:
            pct_class = 'pct-low'

        rows.append({
            'crop': crop,
            'campaing': str(crop.campaing),
            'countryside': str(crop.countryside),
            'lot': str(crop.lot),
            'specie': str(crop.activity.specie),
            'activity': str(crop.activity),
            'ha': ha,
            'est_rate': est_rate,
            'tn_est': tn_est,
            'real_rate': real_rate,
            'tn_real': tn_real,
            'tn_silo': tn_silo,
            'tn_cert': tn_cert,
            'tn_total': tn_total,
            'desvio': desvio,
            'pct': pct,
            'pct_class': pct_class,
            'producer': str(crop.producer) if show_producer else '',
        })

    # ── Totals ─────────────────────────────────────────────────────────
    total_tn_est = sum(r['tn_est'] for r in rows if r['tn_est'] is not None)
    total_tn_real = sum(r['tn_real'] for r in rows if r['tn_real'] is not None)
    total_tn_acopiado = sum(r['tn_total'] for r in rows)
    total_desvio = round(total_tn_real - total_tn_acopiado, 3)

    return {
        'rows': rows,
        'total_tn_est': total_tn_est,
        'total_tn_real': total_tn_real,
        'total_tn_acopiado': total_tn_acopiado,
        'total_desvio': total_desvio,
        'campaigns': Campaing.objects.order_by('-name'),
        'species': Specie.objects.order_by('name'),
        'selected_campaing_id': campaing_id or '',
        'selected_specie_id': specie_id or '',
        'only_with_real': only_with_real,
        'show_producer': show_producer,
    }


def _conciliation_empty_ctx(campaing_id, specie_id, only_with_real, show_producer):
    return {
        'rows': [],
        'total_tn_est': 0, 'total_tn_real': 0,
        'total_tn_acopiado': 0, 'total_desvio': 0,
        'campaigns': Campaing.objects.order_by('-name'),
        'species': Specie.objects.order_by('name'),
        'selected_campaing_id': campaing_id or '',
        'selected_specie_id': specie_id or '',
        'only_with_real': only_with_real,
        'show_producer': show_producer,
    }


@login_required
def conciliation_view(request):
    ctx = _build_conciliation_data(request)
    return render(request, 'storage/report_conciliation.html', {**adminSite.each_context(request), **ctx})


@login_required
def conciliation_excel(request):
    ctx = _build_conciliation_data(request)
    rows = ctx['rows']

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Conciliación'

    hdr_fill = PatternFill('solid', fgColor='1E7D3B')
    hdr_font = Font(bold=True, color='FFFFFF', size=10)
    hdr_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    headers = [
        'Campaña', 'Campo', 'Lote', 'Especie', 'Actividad',
        'Ha cosechadas',
        'Rinde est. (tn/ha)', 'Tn estimadas',
        'Rinde real (tn/ha)', 'Tn reales',
        'Acopiado propio (tn)', 'Acopiado externo (tn)', 'Total acopiado (tn)',
        'Desvío (tn)', '% conciliado',
    ]
    if ctx['show_producer']:
        headers.append('Productor')

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = hdr_align

    num_fmt = '#,##0.000'
    pct_fmt = '#,##0.0"%"'

    for row_idx, r in enumerate(rows, 2):
        ws.cell(row=row_idx, column=1, value=r['campaing'])
        ws.cell(row=row_idx, column=2, value=r['countryside'])
        ws.cell(row=row_idx, column=3, value=r['lot'])
        ws.cell(row=row_idx, column=4, value=r['specie'])
        ws.cell(row=row_idx, column=5, value=r['activity'])
        ws.cell(row=row_idx, column=6, value=r['ha'])
        ws.cell(row=row_idx, column=6).number_format = num_fmt
        ws.cell(row=row_idx, column=7, value=r['est_rate'])
        ws.cell(row=row_idx, column=7).number_format = num_fmt
        ws.cell(row=row_idx, column=8, value=r['tn_est'])
        ws.cell(row=row_idx, column=8).number_format = num_fmt
        ws.cell(row=row_idx, column=9, value=r['real_rate'])
        ws.cell(row=row_idx, column=9).number_format = num_fmt
        ws.cell(row=row_idx, column=10, value=r['tn_real'])
        ws.cell(row=row_idx, column=10).number_format = num_fmt
        ws.cell(row=row_idx, column=11, value=r['tn_silo'])
        ws.cell(row=row_idx, column=11).number_format = num_fmt
        ws.cell(row=row_idx, column=12, value=r['tn_cert'])
        ws.cell(row=row_idx, column=12).number_format = num_fmt
        ws.cell(row=row_idx, column=13, value=r['tn_total'])
        ws.cell(row=row_idx, column=13).number_format = num_fmt
        ws.cell(row=row_idx, column=14, value=r['desvio'])
        ws.cell(row=row_idx, column=14).number_format = num_fmt
        ws.cell(row=row_idx, column=15, value=r['pct'])
        ws.cell(row=row_idx, column=15).number_format = pct_fmt
        if ctx['show_producer']:
            ws.cell(row=row_idx, column=16, value=r['producer'])

    # Totals row
    total_row = len(rows) + 2
    total_fill = PatternFill('solid', fgColor='2c3e50')
    total_font = Font(bold=True, color='FFFFFF', size=10)
    for col in range(1, len(headers) + 1):
        ws.cell(row=total_row, column=col).fill = total_fill
        ws.cell(row=total_row, column=col).font = total_font
    ws.cell(row=total_row, column=1, value='TOTAL')
    ws.cell(row=total_row, column=8, value=ctx['total_tn_est']).number_format = num_fmt
    ws.cell(row=total_row, column=10, value=ctx['total_tn_real']).number_format = num_fmt
    ws.cell(row=total_row, column=13, value=ctx['total_tn_acopiado']).number_format = num_fmt
    ws.cell(row=total_row, column=14, value=ctx['total_desvio']).number_format = num_fmt

    ws.freeze_panes = 'A2'
    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 35)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="conciliacion_rinde.xlsx"'
    wb.save(response)
    return response


# ── Storage Cost Report ──────────────────────────────────────────────────────

def _build_storage_cost_data(request):
    user = request.user
    prod_filter = _producer_filter(user)
    show_producer = user.is_superuser or getattr(user, 'type', None) == 'Administration'

    campaing_id     = request.GET.get('campaing_id')
    provider_id     = request.GET.get('provider_id')
    service_type_id = request.GET.get('service_type_id')
    specie_id       = request.GET.get('specie_id')
    date_from       = request.GET.get('date_from')
    date_to         = request.GET.get('date_to')

    qs = StorageService.objects.select_related(
        'certificate__storage_provider',
        'certificate__campaing',
        'certificate__specie',
        'certificate__producer',
        'type',
        'currency',
    ).order_by(
        'certificate__storage_provider__name',
        'certificate__campaing__name',
        'type__name',
        'currency__name',
        '-date',
    )

    if prod_filter is not None:
        qs = qs.filter(certificate__producer__in=prod_filter)
    if campaing_id:
        qs = qs.filter(certificate__campaing_id=campaing_id)
    if provider_id:
        qs = qs.filter(certificate__storage_provider_id=provider_id)
    if service_type_id:
        qs = qs.filter(type_id=service_type_id)
    if specie_id:
        qs = qs.filter(certificate__specie_id=specie_id)
    if date_from:
        qs = qs.filter(date__gte=date_from)
    if date_to:
        qs = qs.filter(date__lte=date_to)

    services = list(qs)

    # Detail rows
    detail_rows = []
    for s in services:
        cert = s.certificate
        detail_rows.append({
            'date': s.date,
            'provider': cert.storage_provider.name if cert.storage_provider else '—',
            'cert_number': cert.number,
            'campaing': str(cert.campaing),
            'specie': str(cert.specie),
            'service_type': str(s.type),
            'qty_affected': s.quantity_affected,
            'shrinkage_tons': s.shrinkage_tons,
            'cost': s.cost,
            'currency': str(s.currency),
            'notes': s.notes or '',
            'producer': str(cert.producer) if show_producer else '',
        })

    # Summary grouped by (acopiador, campaña, tipo, moneda)
    summary_map = defaultdict(lambda: {'cost': 0.0, 'qty': 0.0, 'count': 0})
    for s in services:
        cert = s.certificate
        key = (
            cert.storage_provider.name if cert.storage_provider else '—',
            str(cert.campaing),
            str(s.type),
            str(s.currency),
        )
        summary_map[key]['cost'] += s.cost
        summary_map[key]['qty']  += s.quantity_affected
        summary_map[key]['count'] += 1

    summary_rows = [
        {
            'provider': k[0], 'campaing': k[1], 'service_type': k[2], 'currency': k[3],
            'total_cost': round(v['cost'], 2),
            'total_qty':  round(v['qty'],  3),
            'count':      v['count'],
        }
        for k, v in sorted(summary_map.items())
    ]

    # Grand totals per currency (can't mix currencies in a single sum)
    totals_by_currency = defaultdict(float)
    total_qty_affected = 0.0
    for s in services:
        totals_by_currency[str(s.currency)] += s.cost
        total_qty_affected += s.quantity_affected
    totals_by_currency = {k: round(v, 2) for k, v in sorted(totals_by_currency.items())}

    return {
        'detail_rows': detail_rows,
        'summary_rows': summary_rows,
        'totals_by_currency': totals_by_currency,
        'total_qty_affected': round(total_qty_affected, 3),
        'show_producer': show_producer,
        'campaigns':     Campaing.objects.order_by('-name'),
        'species':       Specie.objects.order_by('name'),
        'providers':     Provider.objects.filter(type__name='Acopiador').order_by('name'),
        'service_types': StorageServiceType.objects.order_by('name'),
        'selected_campaing_id':     campaing_id     or '',
        'selected_provider_id':     provider_id     or '',
        'selected_service_type_id': service_type_id or '',
        'selected_specie_id':       specie_id       or '',
        'selected_date_from':       date_from       or '',
        'selected_date_to':         date_to         or '',
    }


@login_required
def storage_cost_view(request):
    ctx = _build_storage_cost_data(request)
    return render(request, 'storage/report_storage_cost.html', {**adminSite.each_context(request), **ctx})


@login_required
def storage_cost_excel(request):
    ctx = _build_storage_cost_data(request)

    wb = openpyxl.Workbook()
    hdr_font  = Font(bold=True, color='FFFFFF', size=10)
    hdr_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    num_fmt   = '#,##0.00'
    qty_fmt   = '#,##0.000'

    # ── Sheet 1: Resumen ─────────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = 'Resumen'
    grn_fill = PatternFill('solid', fgColor='1E7D3B')

    sum_headers = ['Acopiador', 'Campaña', 'Tipo de servicio', 'Moneda',
                   'Registros', 'Tn afectadas', 'Costo total']
    for col_idx, h in enumerate(sum_headers, 1):
        cell = ws_sum.cell(row=1, column=col_idx, value=h)
        cell.fill = grn_fill
        cell.font = hdr_font
        cell.alignment = hdr_align

    for row_idx, r in enumerate(ctx['summary_rows'], 2):
        ws_sum.cell(row=row_idx, column=1, value=r['provider'])
        ws_sum.cell(row=row_idx, column=2, value=r['campaing'])
        ws_sum.cell(row=row_idx, column=3, value=r['service_type'])
        ws_sum.cell(row=row_idx, column=4, value=r['currency'])
        ws_sum.cell(row=row_idx, column=5, value=r['count'])
        ws_sum.cell(row=row_idx, column=6, value=r['total_qty']).number_format = qty_fmt
        ws_sum.cell(row=row_idx, column=7, value=r['total_cost']).number_format = num_fmt

    if ctx['totals_by_currency']:
        total_start = len(ctx['summary_rows']) + 3
        total_fill = PatternFill('solid', fgColor='2c3e50')
        total_font = Font(bold=True, color='FFFFFF', size=10)
        for i, (cur, amt) in enumerate(ctx['totals_by_currency'].items()):
            r = total_start + i
            for col in range(1, 8):
                ws_sum.cell(row=r, column=col).fill = total_fill
                ws_sum.cell(row=r, column=col).font = total_font
            ws_sum.cell(row=r, column=1, value=f'TOTAL {cur}')
            ws_sum.cell(row=r, column=7, value=amt).number_format = num_fmt

    ws_sum.freeze_panes = 'A2'
    for col in ws_sum.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=8)
        ws_sum.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 40)

    # ── Sheet 2: Detalle ─────────────────────────────────────────
    ws_det = wb.create_sheet('Detalle')
    blu_fill = PatternFill('solid', fgColor='2c6e9e')

    det_headers = ['Fecha', 'Acopiador', 'Certificado', 'Campaña', 'Especie',
                   'Tipo de servicio', 'Tn afectadas', 'Merma (tn)', 'Costo', 'Moneda', 'Notas']
    if ctx['show_producer']:
        det_headers.append('Productor')

    for col_idx, h in enumerate(det_headers, 1):
        cell = ws_det.cell(row=1, column=col_idx, value=h)
        cell.fill = blu_fill
        cell.font = hdr_font
        cell.alignment = hdr_align

    for row_idx, r in enumerate(ctx['detail_rows'], 2):
        ws_det.cell(row=row_idx, column=1,  value=str(r['date']))
        ws_det.cell(row=row_idx, column=2,  value=r['provider'])
        ws_det.cell(row=row_idx, column=3,  value=r['cert_number'])
        ws_det.cell(row=row_idx, column=4,  value=r['campaing'])
        ws_det.cell(row=row_idx, column=5,  value=r['specie'])
        ws_det.cell(row=row_idx, column=6,  value=r['service_type'])
        ws_det.cell(row=row_idx, column=7,  value=r['qty_affected']).number_format = qty_fmt
        ws_det.cell(row=row_idx, column=8,  value=r['shrinkage_tons']).number_format = qty_fmt
        ws_det.cell(row=row_idx, column=9,  value=r['cost']).number_format = num_fmt
        ws_det.cell(row=row_idx, column=10, value=r['currency'])
        ws_det.cell(row=row_idx, column=11, value=r['notes'])
        if ctx['show_producer']:
            ws_det.cell(row=row_idx, column=12, value=r['producer'])

    ws_det.freeze_panes = 'A2'
    for col in ws_det.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=8)
        ws_det.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 40)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="costos_acopio.xlsx"'
    wb.save(response)
    return response


# ── Stock Aging Report ───────────────────────────────────────────────────────

_AGE_BUCKETS = [
    ('0–30 días',   0,   30,  'age-fresh'),
    ('31–90 días',  31,  90,  'age-moderate'),
    ('91–180 días', 91,  180, 'age-old'),
    ('>180 días',   181, None,'age-very-old'),
]


def _bucket_for(age_days):
    for label, lo, hi, css in _AGE_BUCKETS:
        if hi is None or age_days <= hi:
            return label, css
    return _AGE_BUCKETS[-1][0], _AGE_BUCKETS[-1][3]


def _build_aging_data(request):
    user = request.user
    prod_filter = _producer_filter(user)
    show_producer = user.is_superuser or getattr(user, 'type', None) == 'Administration'

    campaing_id = request.GET.get('campaing_id')
    provider_id = request.GET.get('provider_id')
    specie_id   = request.GET.get('specie_id')
    bucket_filter = request.GET.get('bucket', '')    # label text to narrow display

    today = date.today()

    qs = StorageCertificate.objects.select_related(
        'storage_provider', 'campaing', 'specie', 'producer',
    ).annotate(
        tons_in=Coalesce(
            Sum('movements__quantity_tons',
                filter=Q(movements__type='entry')), 0.0),
        tons_out=Coalesce(
            Sum('movements__quantity_tons',
                filter=Q(movements__type__in=['sale_exit', 'withdrawal', 'transfer'])), 0.0),
    ).filter(
        state__in=[
            StorageCertificate.ACTIVE,
            StorageCertificate.PARTIAL,
            StorageCertificate.WARRANT,
        ]
    ).order_by('date_issue')  # oldest first

    if prod_filter is not None:
        qs = qs.filter(producer__in=prod_filter)
    if campaing_id:
        qs = qs.filter(campaing_id=campaing_id)
    if provider_id:
        qs = qs.filter(storage_provider_id=provider_id)
    if specie_id:
        qs = qs.filter(specie_id=specie_id)

    rows = []
    for cert in qs:
        qty_avail = max(cert.tons_in - cert.tons_out, 0)
        if qty_avail <= 0:
            continue  # cert is technically empty despite active state
        age_days = (today - cert.date_issue).days if cert.date_issue else None
        pct_remaining = (
            round(qty_avail / cert.quantity_entry * 100, 1)
            if cert.quantity_entry else None
        )
        days_to_expiry = (
            (cert.date_expiry - today).days if cert.date_expiry else None
        )
        bucket_label, age_css = _bucket_for(age_days) if age_days is not None else ('—', '')

        rows.append({
            'cert': cert,
            'number': cert.number,
            'provider': cert.storage_provider.name if cert.storage_provider else '—',
            'campaing': str(cert.campaing),
            'specie': str(cert.specie),
            'date_issue': cert.date_issue,
            'date_expiry': cert.date_expiry,
            'days_to_expiry': days_to_expiry,
            'age_days': age_days,
            'age_class': _age_class(age_days),
            'bucket': bucket_label,
            'qty_entry': cert.quantity_entry,
            'qty_available': round(qty_avail, 3),
            'pct_remaining': pct_remaining,
            'state': cert.get_state_display(),
            'is_warrant': cert.is_warrant,
            'producer': str(cert.producer) if show_producer else '',
        })

    if bucket_filter:
        rows = [r for r in rows if r['bucket'] == bucket_filter]

    # Bucket summary
    bucket_map = {b[0]: {'count': 0, 'tons': 0.0, 'css': b[3]} for b in _AGE_BUCKETS}
    for r in rows:
        bl = r['bucket']
        if bl in bucket_map:
            bucket_map[bl]['count'] += 1
            bucket_map[bl]['tons']  += r['qty_available']
    total_tons = sum(v['tons'] for v in bucket_map.values())
    bucket_summary = [
        {
            'label': b[0], 'count': bucket_map[b[0]]['count'],
            'tons': round(bucket_map[b[0]]['tons'], 3),
            'pct_tons': round(bucket_map[b[0]]['tons'] / total_tons * 100, 1) if total_tons else 0,
            'css': b[3],
        }
        for b in _AGE_BUCKETS
    ]

    avg_days = (
        round(sum(r['age_days'] for r in rows if r['age_days'] is not None) / len(rows), 0)
        if rows else None
    )

    return {
        'rows': rows,
        'bucket_summary': bucket_summary,
        'total_certs': len(rows),
        'total_tons': round(total_tons, 3),
        'avg_days': avg_days,
        'show_producer': show_producer,
        'campaigns':  Campaing.objects.order_by('-name'),
        'species':    Specie.objects.order_by('name'),
        'providers':  Provider.objects.filter(type__name='Acopiador').order_by('name'),
        'age_buckets': [b[0] for b in _AGE_BUCKETS],
        'selected_campaing_id': campaing_id  or '',
        'selected_provider_id': provider_id  or '',
        'selected_specie_id':   specie_id    or '',
        'selected_bucket':      bucket_filter or '',
    }


@login_required
def aging_view(request):
    ctx = _build_aging_data(request)
    return render(request, 'storage/report_aging.html', {**adminSite.each_context(request), **ctx})


@login_required
def aging_excel(request):
    ctx = _build_aging_data(request)

    wb = openpyxl.Workbook()
    hdr_font  = Font(bold=True, color='FFFFFF', size=10)
    hdr_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    num_fmt = '#,##0.000'

    # ── Sheet 1: Resumen por tramo ────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = 'Resumen por tramo'
    grn_fill = PatternFill('solid', fgColor='1E7D3B')

    sum_headers = ['Tramo de antigüedad', 'Certificados', 'Tn disponibles', '% del total']
    for col_idx, h in enumerate(sum_headers, 1):
        cell = ws_sum.cell(row=1, column=col_idx, value=h)
        cell.fill = grn_fill; cell.font = hdr_font; cell.alignment = hdr_align

    for row_idx, b in enumerate(ctx['bucket_summary'], 2):
        ws_sum.cell(row=row_idx, column=1, value=b['label'])
        ws_sum.cell(row=row_idx, column=2, value=b['count'])
        ws_sum.cell(row=row_idx, column=3, value=b['tons']).number_format = num_fmt
        ws_sum.cell(row=row_idx, column=4, value=b['pct_tons'])

    total_row = len(ctx['bucket_summary']) + 2
    total_fill = PatternFill('solid', fgColor='2c3e50')
    total_font = Font(bold=True, color='FFFFFF', size=10)
    for col in range(1, 5):
        ws_sum.cell(row=total_row, column=col).fill = total_fill
        ws_sum.cell(row=total_row, column=col).font = total_font
    ws_sum.cell(row=total_row, column=1, value='TOTAL')
    ws_sum.cell(row=total_row, column=2, value=ctx['total_certs'])
    ws_sum.cell(row=total_row, column=3, value=ctx['total_tons']).number_format = num_fmt

    ws_sum.freeze_panes = 'A2'
    for col in ws_sum.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=8)
        ws_sum.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 35)

    # ── Sheet 2: Detalle ─────────────────────────────────────────
    ws_det = wb.create_sheet('Detalle')
    blu_fill = PatternFill('solid', fgColor='2c6e9e')

    det_headers = ['Certificado', 'Acopiador', 'Campaña', 'Especie',
                   'Fecha emisión', 'Días', 'Tramo',
                   'Tn entrada', 'Tn disponibles', '% restante',
                   'Vencimiento', 'Días p/vencer', 'Estado']
    if ctx['show_producer']:
        det_headers.append('Productor')

    for col_idx, h in enumerate(det_headers, 1):
        cell = ws_det.cell(row=1, column=col_idx, value=h)
        cell.fill = blu_fill; cell.font = hdr_font; cell.alignment = hdr_align

    for row_idx, r in enumerate(ctx['rows'], 2):
        ws_det.cell(row=row_idx, column=1,  value=r['number'])
        ws_det.cell(row=row_idx, column=2,  value=r['provider'])
        ws_det.cell(row=row_idx, column=3,  value=r['campaing'])
        ws_det.cell(row=row_idx, column=4,  value=r['specie'])
        ws_det.cell(row=row_idx, column=5,  value=str(r['date_issue']))
        ws_det.cell(row=row_idx, column=6,  value=r['age_days'])
        ws_det.cell(row=row_idx, column=7,  value=r['bucket'])
        ws_det.cell(row=row_idx, column=8,  value=r['qty_entry']).number_format = num_fmt
        ws_det.cell(row=row_idx, column=9,  value=r['qty_available']).number_format = num_fmt
        ws_det.cell(row=row_idx, column=10, value=r['pct_remaining'])
        ws_det.cell(row=row_idx, column=11, value=str(r['date_expiry']) if r['date_expiry'] else '')
        ws_det.cell(row=row_idx, column=12, value=r['days_to_expiry'])
        ws_det.cell(row=row_idx, column=13, value=r['state'])
        if ctx['show_producer']:
            ws_det.cell(row=row_idx, column=14, value=r['producer'])

    ws_det.freeze_panes = 'A2'
    for col in ws_det.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=8)
        ws_det.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 35)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="antiguedad_stock_terceros.xlsx"'
    wb.save(response)
    return response
