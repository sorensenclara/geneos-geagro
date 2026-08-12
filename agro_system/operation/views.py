from user.admin.adminSite import adminSite
from django.shortcuts import render
from operation.models.work_order import WorkOrder
from operation.models.input_work_order import InputWorkOrder
from django.http import HttpResponse, HttpResponseRedirect
from xhtml2pdf import pisa
from django.template.loader import get_template
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse
from django.contrib import messages
from django.db.models import Sum, F, ExpressionWrapper, FloatField
from django.db.models.functions import TruncMonth, Coalesce
from collections import defaultdict
from parameter.models.campaign import Campaing
from parameter.models.ot_state import OTState
from parameter.models.area import Area
from input.models.purchase_order_detail import PurchaseOrderDetail
from user.models.user_producer import UserProducer
from parameter.models.currency import Currency
from operation.models.crop import Crop
from producer.models.contractor import Contractor


@login_required
def WorkOrderPDFView(request):

    ids = request.GET.get('ids')
    ids_list = ids.split(',') if ids else []
    orders = WorkOrder.objects.filter(id__in=ids_list)
    producer_filter = _producer_filter(request.user)
    if producer_filter is not None:
        orders = orders.filter(crop__producer__in=producer_filter)
    if orders:
        # Renderiza la plantilla HTML y pasa el queryset como contexto
        context = {'orders': orders}

        template_path = 'operations/report_work_order.html'
        # Create a Django response object, and specify content_type as pdf
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="reporte-ordenes-trabajo.pdf"'
        # find the template and render it.
        template = get_template(template_path)
        html = template.render(context)

        # create a pdf
        pisa_status = pisa.CreatePDF(
        html, dest=response)
        # if error then show some funny view
        if pisa_status.err:
            return HttpResponse('Errores: <pre>' + html + '</pre>')
        return response

    else:
        messages.error(request, "No hay pendientes para generar el reporte.")
        return HttpResponseRedirect(request.META.get('HTTP_REFERER', '/'))


@login_required
def calendar_json(request):
    producer_filter = _producer_filter(request.user)
    orders = WorkOrder.objects.all()
    if producer_filter is not None:
        orders = orders.filter(crop__producer__in=producer_filter)
    data = []
    for order in orders:
        data.append({
            'title': f"{order.code} ({order.producer} - {order.campaing})",
            'start': order.init_order.isoformat(),
            'url': f'/admin/operation/workorder/{order.id}/change/',  # Opcional: enlace al admin
        })
    return JsonResponse(data, safe=False)


@login_required
def calendar_view(request):
    return render(request, 'operations/calendar.html', adminSite.each_context(request))


def _producer_filter(user):
    if user.is_superuser or getattr(user, 'type', None) == 'Administration':
        return None
    if getattr(user, 'type', None) == 'Organization':
        return UserProducer.objects.filter(association=user)
    if getattr(user, 'type', None) == 'Producer':
        return UserProducer.objects.filter(email=user.email)
    if getattr(user, 'type', None) == 'Staff':
        from producer.models.staff import Staff
        staff = Staff.objects.filter(user=user).first()
        if staff:
            return UserProducer.objects.filter(pk=staff.producer_id)
    return UserProducer.objects.none()


def _build_input_report_data(request):
    """Construye todos los datos del informe de insumos. Compartido entre HTML y Excel."""
    user = request.user
    campaing_id = request.GET.get('campaing_id')
    use_base = request.GET.get('use_base') == '1'
    base_currency = Currency.objects.filter(is_base=True).first()
    producer_filter = _producer_filter(user)

    context = {
        'campaigns': Campaing.objects.all().order_by('-date_init'),
        'selected_campaing': None,
        'use_base': use_base,
        'base_currency': base_currency,
        'purchases': [],
        'purchases_groups': [],
        'pmp_list': [],
        'usages': [],
    }

    if not campaing_id:
        return context

    try:
        campaign = Campaing.objects.get(pk=campaing_id)
        context['selected_campaing'] = campaign

        purchase_qs = PurchaseOrderDetail.objects.filter(
            purchase_order__state='confirm',
            purchase_order__date__gte=campaign.date_init,
            purchase_order__date__lte=campaign.date_final,
        ).select_related(
            'input', 'input__unit_measure',
            'purchase_order', 'purchase_order__provider', 'purchase_order__currency',
        )
        if producer_filter is not None:
            purchase_qs = purchase_qs.filter(purchase_order__producer__in=producer_filter)

        if use_base:
            purchases = list(purchase_qs.order_by('input__name', 'purchase_order__date'))
            context['purchases'] = purchases
            context['purchases_total'] = sum(
                (p.total_base if p.total_base is not None else (p.total or 0)) for p in purchases
            )
            pmp_map = defaultdict(lambda: {'name': '', 'unit': '-', 'qty': 0.0, 'cost': 0.0})
            for p in purchases:
                if p.amount:
                    d = pmp_map[p.input_id]
                    d['name'] = p.input.name
                    d['unit'] = p.input.unit_measure.name if p.input.unit_measure else '-'
                    d['qty'] += p.amount
                    d['cost'] += p.total_base if p.total_base is not None else (p.total or 0)
            context['pmp_list'] = sorted(
                [{**d, 'pmp': d['cost'] / d['qty'] if d['qty'] else 0} for d in pmp_map.values()],
                key=lambda x: x['name'],
            )
        else:
            groups_map = {}
            groups_order = []
            for p in purchase_qs.order_by('purchase_order__currency__name', 'input__name', 'purchase_order__date'):
                currency_obj = p.purchase_order.currency
                key = currency_obj.pk if currency_obj else '__none__'
                if key not in groups_map:
                    groups_map[key] = {
                        'currency': currency_obj,
                        'currency_name': str(currency_obj) if currency_obj else 'Sin moneda',
                        'items': [],
                        'total': 0.0,
                        '_pmp': defaultdict(lambda: {'name': '', 'unit': '-', 'qty': 0.0, 'cost': 0.0}),
                    }
                    groups_order.append(key)
                g = groups_map[key]
                g['items'].append(p)
                g['total'] += p.total or 0
                if p.price_unit and p.amount:
                    d = g['_pmp'][p.input_id]
                    d['name'] = p.input.name
                    d['unit'] = p.input.unit_measure.name if p.input.unit_measure else '-'
                    d['qty'] += p.amount
                    d['cost'] += p.amount * p.price_unit
            purchases_groups = []
            for key in groups_order:
                g = groups_map[key]
                purchases_groups.append({
                    'currency': g['currency'],
                    'currency_name': g['currency_name'],
                    'items': g['items'],
                    'total': g['total'],
                    'pmp_list': sorted(
                        [{**d, 'pmp': d['cost'] / d['qty'] if d['qty'] else 0} for d in g['_pmp'].values()],
                        key=lambda x: x['name'],
                    ),
                })
            context['purchases_groups'] = purchases_groups

        usage_qs = InputWorkOrder.objects.filter(
            work_order__campaing_id=campaing_id,
        ).select_related('input', 'input__unit_measure', 'input__currency', 'work_order', 'work_order__stage', 'work_order__crop')
        if producer_filter is not None:
            usage_qs = usage_qs.filter(work_order__producer__in=producer_filter)

        if use_base:
            usages_converted = []
            for item in usage_qs.order_by('input__name', 'work_order__date_init'):
                real_dose = item.dose_applied if item.dose_applied is not None else item.dose
                usages_converted.append({
                    'obj': item,
                    'value_estimated_base': item.value_estimated,
                    'total_estimated_base': (item.value_estimated or 0) * (item.dose or 0) or None,
                    'value_real_base': item.value_real_base,
                    'total_real_base': (item.value_real_base or 0) * (real_dose or 0) or None,
                    'rate': item.exchange_rate or 1.0,
                })
            context['usages'] = usages_converted
        else:
            usages_list = []
            for item in usage_qs.order_by('input__name', 'work_order__date_init'):
                real_dose = item.dose_applied if item.dose_applied is not None else item.dose
                usages_list.append({
                    'obj': item,
                    'value_estimated': item.value_estimated,
                    'total_estimated': (item.value_estimated or 0) * (item.dose or 0) or None,
                    'value_real': item.value_real,
                    'total_real': (item.value_real or 0) * (real_dose or 0) or None,
                })
            context['usages'] = usages_list

    except Campaing.DoesNotExist:
        pass

    return context


def _build_budget_data(request):
    """Construye todos los datos del informe presupuesto vs real. Compartido entre HTML y Excel."""
    user = request.user
    campaing_id = request.GET.get('campaing_id')
    use_base = request.GET.get('use_base') == '1'
    base_currency = Currency.objects.filter(is_base=True).first()
    base_symbol = (base_currency.symbol or '$') if base_currency else '$'
    producer_filter = _producer_filter(user)
    state_id = request.GET.get('state_id') or ''

    context = {
        'campaigns': Campaing.objects.all().order_by('-date_init'),
        'selected_campaing': None,
        'all_states': OTState.objects.all().order_by('name'),
        'selected_state_id': state_id,
        'use_base': use_base,
        'base_currency': base_currency,
        'base_symbol': base_symbol,
        'report_rows': [],
        'months': [],
        'currency_totals': [],
    }

    if not campaing_id:
        return context

    try:
        campaign = Campaing.objects.get(pk=campaing_id)
        context['selected_campaing'] = campaign

        months_set = set()
        crop_map = defaultdict(lambda: defaultdict(lambda: {'presupuestado': 0.0, 'real': 0.0}))

        if use_base:
            usage_items = InputWorkOrder.objects.filter(
                work_order__campaing_id=campaing_id,
            ).select_related('work_order__crop')
            if producer_filter is not None:
                usage_items = usage_items.filter(work_order__producer__in=producer_filter)
            if state_id:
                usage_items = usage_items.filter(work_order__stateOT_id=state_id)

            for item in usage_items:
                if not item.work_order.date_init:
                    continue
                month = item.work_order.date_init.replace(day=1)
                months_set.add(month)
                crop = item.work_order.crop.name if item.work_order.crop else '(sin cultivo)'
                # value_estimated ya está en moneda base
                est = (item.value_estimated or 0) * (item.dose or 0)
                # value_real_base ya tiene la conversión histórica; si no existe, cae a value_estimated
                real_base = item.value_real_base if item.value_real_base is not None else item.value_estimated
                real_dose = item.dose_applied if item.dose_applied is not None else item.dose
                real = (real_base or 0) * (real_dose or 0)
                crop_map[crop][month]['presupuestado'] += est
                crop_map[crop][month]['real'] += real

            # Incluir costo de contratista por OT (en moneda base)
            wo_qs = WorkOrder.objects.filter(
                campaing_id=campaing_id,
            ).exclude(
                contractor_cost__isnull=True, contractor_cost_real__isnull=True,
            ).select_related('crop')
            if producer_filter is not None:
                wo_qs = wo_qs.filter(producer__in=producer_filter)
            if state_id:
                wo_qs = wo_qs.filter(stateOT_id=state_id)
            for wo in wo_qs:
                if not wo.date_init:
                    continue
                month = wo.date_init.replace(day=1)
                months_set.add(month)
                crop = wo.crop.name if wo.crop else '(sin cultivo)'
                if wo.contractor_cost:
                    crop_map[crop][month]['presupuestado'] += wo.contractor_cost
                real_cost = wo.contractor_cost_real if wo.contractor_cost_real is not None else wo.contractor_cost
                if real_cost:
                    crop_map[crop][month]['real'] += real_cost
        else:
            usage_items = InputWorkOrder.objects.filter(
                work_order__campaing_id=campaing_id,
            ).select_related('currency', 'input__currency', 'work_order__crop')
            if producer_filter is not None:
                usage_items = usage_items.filter(work_order__producer__in=producer_filter)
            if state_id:
                usage_items = usage_items.filter(work_order__stateOT_id=state_id)

            currency_symbols = {}  # cache: currency_name → symbol
            for item in usage_items:
                if not item.work_order.date_init:
                    continue
                # La moneda del real es item.currency; si no está, la del insumo
                currency = item.currency or (item.input.currency if item.input else None)
                currency_name = str(currency) if currency else 'Sin moneda'
                if currency_name not in currency_symbols:
                    currency_symbols[currency_name] = (currency.symbol or currency_name) if currency else currency_name
                month = item.work_order.date_init.replace(day=1)
                months_set.add(month)
                crop = item.work_order.crop.name if item.work_order.crop else '(sin cultivo)'
                crop_key = f"[{currency_name}] {crop}"
                # value_estimated está en moneda base → se convierte a la moneda real dividiendo por la tasa.
                # Así pres y real quedan en la misma moneda y el desvío es comparable.
                # Para ítems en moneda base, exchange_rate=1 y la división no cambia nada.
                rate = item.exchange_rate or 1.0
                est = (item.value_estimated or 0) * (item.dose or 0) / rate
                real_val = item.value_real if item.value_real is not None else None
                real_dose = item.dose_applied if item.dose_applied is not None else item.dose
                real = (real_val or 0) * (real_dose or 0)
                crop_map[crop_key][month]['presupuestado'] += est
                crop_map[crop_key][month]['real'] += real

            # Incluir costo de contratista (siempre en moneda base)
            base_currency_name = str(base_currency) if base_currency else 'Sin moneda'
            wo_qs = WorkOrder.objects.filter(
                campaing_id=campaing_id,
            ).exclude(
                contractor_cost__isnull=True, contractor_cost_real__isnull=True,
            ).select_related('crop')
            if producer_filter is not None:
                wo_qs = wo_qs.filter(producer__in=producer_filter)
            if state_id:
                wo_qs = wo_qs.filter(stateOT_id=state_id)
            for wo in wo_qs:
                if not wo.date_init:
                    continue
                month = wo.date_init.replace(day=1)
                months_set.add(month)
                crop = wo.crop.name if wo.crop else '(sin cultivo)'
                crop_key = f"[{base_currency_name}] {crop}"
                if wo.contractor_cost:
                    crop_map[crop_key][month]['presupuestado'] += wo.contractor_cost
                real_cost = wo.contractor_cost_real if wo.contractor_cost_real is not None else wo.contractor_cost
                if real_cost:
                    crop_map[crop_key][month]['real'] += real_cost

        months = sorted(months_set)
        context['months'] = months

        def _make_row(crop_name, month_data):
            total_pres = sum(v['presupuestado'] for v in month_data.values())
            total_real = sum(v['real'] for v in month_data.values())
            if total_pres == 0 and total_real == 0:
                return None
            desvio = total_real - total_pres
            desvio_pct = (desvio / total_pres * 100) if total_pres else None
            cells = [month_data.get(m, {'presupuestado': 0.0, 'real': 0.0}) for m in months]
            max_val = max(total_pres, total_real)
            return {
                'crop': crop_name,
                'cells': cells,
                'total_pres': total_pres,
                'total_real': total_real,
                'desvio': desvio,
                'desvio_pct': desvio_pct,
                'bar_pres': round(total_pres / max_val * 100) if max_val else 0,
                'bar_real': round(total_real / max_val * 100) if max_val else 0,
            }

        report_rows = [r for r in (
            _make_row(cn, md) for cn, md in sorted(crop_map.items())
        ) if r is not None]
        context['report_rows'] = report_rows

        if use_base:
            month_totals = []
            for i in range(len(months)):
                pres = sum(r['cells'][i]['presupuestado'] for r in report_rows)
                real = sum(r['cells'][i]['real'] for r in report_rows)
                month_totals.append({'presupuestado': pres, 'real': real})
            context['month_totals'] = month_totals
            context['grand_pres'] = sum(r['total_pres'] for r in report_rows)
            context['grand_real'] = sum(r['total_real'] for r in report_rows)
            grand_desvio = context['grand_real'] - context['grand_pres']
            context['grand_desvio'] = grand_desvio
            context['grand_desvio_pct'] = (grand_desvio / context['grand_pres'] * 100) if context['grand_pres'] else None
        else:
            grouped = {}
            for row in report_rows:
                crop_key = row['crop']
                cn = crop_key[1:crop_key.index(']')] if crop_key.startswith('[') else 'Sin moneda'
                crop_clean = crop_key[crop_key.index(']') + 2:] if crop_key.startswith('[') else crop_key
                if cn not in grouped:
                    grouped[cn] = {
                        'currency': cn,
                        'symbol': currency_symbols.get(cn, cn),
                        'rows': [],
                        'month_cells': [{'presupuestado': 0.0, 'real': 0.0} for _ in months],
                        'total_pres': 0.0,
                        'total_real': 0.0,
                    }
                g = grouped[cn]
                g['rows'].append({**row, 'crop': crop_clean})
                g['total_pres'] += row['total_pres']
                g['total_real'] += row['total_real']
                for i, cell in enumerate(row['cells']):
                    g['month_cells'][i]['presupuestado'] += cell['presupuestado']
                    g['month_cells'][i]['real'] += cell['real']

            grouped_by_currency = []
            for cn, g in sorted(grouped.items()):
                desvio = g['total_real'] - g['total_pres']
                desvio_pct = (desvio / g['total_pres'] * 100) if g['total_pres'] else None
                max_val = max(g['total_pres'], g['total_real'])
                grouped_by_currency.append({
                    **g,
                    'desvio': desvio,
                    'desvio_pct': desvio_pct,
                    'bar_pres': round(g['total_pres'] / max_val * 100) if max_val else 0,
                    'bar_real': round(g['total_real'] / max_val * 100) if max_val else 0,
                })

            context['grouped_by_currency'] = grouped_by_currency
            # Mantener currency_totals para el Excel (usa misma estructura)
            context['currency_totals'] = grouped_by_currency

    except Campaing.DoesNotExist:
        pass

    return context


@login_required
def input_report_view(request):
    context = _build_input_report_data(request)
    return render(request, 'operations/report_inputs.html', {**adminSite.each_context(request), **context})


@login_required
def input_report_excel(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    ctx = _build_input_report_data(request)

    wb = openpyxl.Workbook()
    GREEN = '1E7D3B'
    BLUE = '2c6e9e'
    HDR_FONT = Font(bold=True, color='FFFFFF')
    HDR_FILL_G = PatternFill('solid', fgColor=GREEN)
    HDR_FILL_B = PatternFill('solid', fgColor=BLUE)
    SUB_FILL = PatternFill('solid', fgColor='E8F4E8')
    CENTER = Alignment(horizontal='center')
    RIGHT = Alignment(horizontal='right')

    def set_header(ws, row, cols, fill):
        for c, val in enumerate(cols, 1):
            cell = ws.cell(row=row, column=c, value=val)
            cell.font = HDR_FONT
            cell.fill = fill
            cell.alignment = CENTER

    def autowidth(ws):
        for col in ws.columns:
            max_len = max((len(str(c.value or '')) for c in col), default=10)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)

    campaign_name = ctx['selected_campaing'].name if ctx['selected_campaing'] else 'sin_campaña'
    use_base = ctx['use_base']
    base_label = str(ctx['base_currency']) if ctx['base_currency'] else 'base'

    # ── Hoja 1: Compras ───────────────────────────────────────
    ws1 = wb.active
    ws1.title = 'Compras'
    if use_base:
        set_header(ws1, 1, ['Insumo', 'Proveedor', 'Fecha OC', 'N° OC', 'Moneda OC',
                             'Cantidad', 'Unidad', f'Total original', f'Total {base_label}'], HDR_FILL_G)
        for r, p in enumerate(ctx['purchases'], 2):
            ws1.cell(r, 1, p.input.name)
            ws1.cell(r, 2, str(p.purchase_order.provider))
            ws1.cell(r, 3, str(p.purchase_order.date))
            ws1.cell(r, 4, f"OC{p.purchase_order.number:04d}")
            ws1.cell(r, 5, str(p.purchase_order.currency) if p.purchase_order.currency else '')
            ws1.cell(r, 6, p.amount)
            ws1.cell(r, 7, p.input.unit_measure.name if p.input.unit_measure else '')
            ws1.cell(r, 8, p.total or '')
            ws1.cell(r, 9, p.total_base if p.total_base is not None else (p.total or ''))
    else:
        set_header(ws1, 1, ['Moneda', 'Insumo', 'Proveedor', 'Fecha OC', 'N° OC',
                             'Cantidad', 'Unidad', 'Precio unitario', 'Impuesto', 'Total'], HDR_FILL_G)
        r = 2
        for group in ctx['purchases_groups']:
            for p in group['items']:
                ws1.cell(r, 1, group['currency_name'])
                ws1.cell(r, 2, p.input.name)
                ws1.cell(r, 3, str(p.purchase_order.provider))
                ws1.cell(r, 4, str(p.purchase_order.date))
                ws1.cell(r, 5, f"OC{p.purchase_order.number:04d}")
                ws1.cell(r, 6, p.amount)
                ws1.cell(r, 7, p.input.unit_measure.name if p.input.unit_measure else '')
                ws1.cell(r, 8, p.price_unit or '')
                ws1.cell(r, 9, p.tax or '')
                ws1.cell(r, 10, p.total or '')
                r += 1
    autowidth(ws1)

    # ── Hoja 2: PMP ───────────────────────────────────────────
    ws2 = wb.create_sheet('PMP')
    if use_base:
        set_header(ws2, 1, ['Insumo', 'Cantidad total', 'Unidad', f'Costo total {base_label}', f'PMP en {base_label}'], HDR_FILL_G)
        for r, p in enumerate(ctx['pmp_list'], 2):
            ws2.cell(r, 1, p['name'])
            ws2.cell(r, 2, p['qty'])
            ws2.cell(r, 3, p['unit'])
            ws2.cell(r, 4, p['cost'])
            ws2.cell(r, 5, p['pmp'])
    else:
        set_header(ws2, 1, ['Moneda', 'Insumo', 'Cantidad total', 'Unidad', 'Costo total', 'PMP'], HDR_FILL_G)
        r = 2
        for group in ctx['purchases_groups']:
            for p in group['pmp_list']:
                ws2.cell(r, 1, group['currency_name'])
                ws2.cell(r, 2, p['name'])
                ws2.cell(r, 3, p['qty'])
                ws2.cell(r, 4, p['unit'])
                ws2.cell(r, 5, p['cost'])
                ws2.cell(r, 6, p['pmp'])
                r += 1
    autowidth(ws2)

    # ── Hoja 3: Uso OTs ───────────────────────────────────────
    ws3 = wb.create_sheet('Uso OTs')
    if use_base:
        set_header(ws3, 1, ['Insumo', 'Moneda insumo', 'OT', 'Cultivo', 'Dosis plan.',
                             'Dosis aplic.', 'Unidad',
                             f'P.unit est. {base_label}', f'Total est. {base_label}',
                             f'P.unit real {base_label}', f'Total real {base_label}'], HDR_FILL_B)
        for r, u in enumerate(ctx['usages'], 2):
            item = u['obj']
            ws3.cell(r, 1, item.input.name)
            ws3.cell(r, 2, str(item.input.currency) if item.input.currency else '')
            ws3.cell(r, 3, item.work_order.code)
            ws3.cell(r, 4, str(item.work_order.crop) if item.work_order.crop else '')
            ws3.cell(r, 5, item.dose)
            ws3.cell(r, 6, item.dose_applied or '')
            ws3.cell(r, 7, item.input.unit_measure.name if item.input.unit_measure else '')
            ws3.cell(r, 8, u['value_estimated_base'] if u['value_estimated_base'] is not None else '')
            ws3.cell(r, 9, u['total_estimated_base'] if u['total_estimated_base'] else '')
            ws3.cell(r, 10, u['value_real_base'] if u['value_real_base'] is not None else '')
            ws3.cell(r, 11, u['total_real_base'] if u['total_real_base'] else '')
    else:
        set_header(ws3, 1, ['Insumo', 'Moneda', 'OT', 'Cultivo', 'Dosis plan.',
                             'Dosis aplic.', 'Unidad',
                             'P.unit estimado', 'Total estimado',
                             'P.unit real', 'Total real'], HDR_FILL_B)
        for r, u in enumerate(ctx['usages'], 2):
            item = u['obj']
            ws3.cell(r, 1, item.input.name)
            ws3.cell(r, 2, str(item.input.currency) if item.input.currency else '')
            ws3.cell(r, 3, item.work_order.code)
            ws3.cell(r, 4, str(item.work_order.crop) if item.work_order.crop else '')
            ws3.cell(r, 5, item.dose)
            ws3.cell(r, 6, item.dose_applied or '')
            ws3.cell(r, 7, item.input.unit_measure.name if item.input.unit_measure else '')
            ws3.cell(r, 8, u['value_estimated'] or '')
            ws3.cell(r, 9, u['total_estimated'] or '')
            ws3.cell(r, 10, u['value_real'] or '')
            ws3.cell(r, 11, u['total_real'] or '')
    autowidth(ws3)

    resp = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="insumos_{campaign_name}.xlsx"'
    wb.save(resp)
    return resp


@login_required
def budget_vs_real_view(request):
    context = _build_budget_data(request)
    return render(request, 'operations/report_budget_vs_real.html', {**adminSite.each_context(request), **context})


@login_required
def budget_vs_real_excel(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    ctx = _build_budget_data(request)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Pres vs Real'

    HDR_FONT = Font(bold=True, color='FFFFFF')
    FILL_DARK = PatternFill('solid', fgColor='2C3E50')
    FILL_GREEN = PatternFill('solid', fgColor='1E7D3B')
    FILL_RED = PatternFill('solid', fgColor='8B1A1A')
    FILL_MONTH = PatternFill('solid', fgColor='E8F0E8')

    months = ctx['months']
    report_rows = ctx['report_rows']
    use_base = ctx['use_base']
    campaign_name = ctx['selected_campaing'].name if ctx['selected_campaing'] else 'sin_campaña'
    base_label = str(ctx['base_currency']) if ctx['base_currency'] else 'base'

    # Fila 1: encabezados de mes
    ws.cell(1, 1, 'Cultivo').font = Font(bold=True)
    col = 2
    for m in months:
        cell = ws.cell(1, col, m.strftime('%b %Y'))
        cell.font = HDR_FONT
        cell.fill = FILL_MONTH
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 1)
        col += 2

    for label, fill in [('Total campaña', FILL_GREEN), ('Desvío', FILL_RED)]:
        cell = ws.cell(1, col, label)
        cell.font = HDR_FONT
        cell.fill = fill
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 1)
        col += 2

    # Fila 2: sub-headers
    ws.cell(2, 1, '')
    col = 2
    for _ in months:
        ws.cell(2, col, 'Pres.').font = Font(bold=True)
        ws.cell(2, col + 1, 'Real').font = Font(bold=True)
        col += 2
    for lbl in ['Pres.', 'Real', '$ Desvío', '% Desvío']:
        ws.cell(2, col, lbl).font = Font(bold=True)
        col += 1

    # Filas de datos
    for i, row in enumerate(report_rows, 3):
        ws.cell(i, 1, row['crop'])
        col = 2
        for cell_data in row['cells']:
            ws.cell(i, col, round(cell_data['presupuestado'], 0) or '')
            ws.cell(i, col + 1, round(cell_data['real'], 0) or '')
            col += 2
        ws.cell(i, col, round(row['total_pres'], 0))
        ws.cell(i, col + 1, round(row['total_real'], 0))
        ws.cell(i, col + 2, round(row['desvio'], 0))
        ws.cell(i, col + 3, f"{row['desvio_pct']:.1f}%" if row['desvio_pct'] is not None else '')

    # Pie: totales
    next_row = len(report_rows) + 3
    if use_base:
        mt = ctx.get('month_totals', [])
        ws.cell(next_row, 1, 'TOTAL').font = Font(bold=True)
        col = 2
        for m_total in mt:
            ws.cell(next_row, col, round(m_total['presupuestado'], 0) or '')
            ws.cell(next_row, col + 1, round(m_total['real'], 0) or '')
            col += 2
        ws.cell(next_row, col, round(ctx.get('grand_pres', 0), 0)).font = Font(bold=True)
        ws.cell(next_row, col + 1, round(ctx.get('grand_real', 0), 0)).font = Font(bold=True)
    else:
        for ct in ctx.get('currency_totals', []):
            ws.cell(next_row, 1, f"TOTAL {ct['currency']}").font = Font(bold=True)
            col = 2
            for mc in ct['month_cells']:
                ws.cell(next_row, col, round(mc['presupuestado'], 0) or '')
                ws.cell(next_row, col + 1, round(mc['real'], 0) or '')
                col += 2
            ws.cell(next_row, col, round(ct['total_pres'], 0)).font = Font(bold=True)
            ws.cell(next_row, col + 1, round(ct['total_real'], 0)).font = Font(bold=True)
            ws.cell(next_row, col + 2, round(ct['desvio'], 0)).font = Font(bold=True)
            ws.cell(next_row, col + 3, f"{ct['desvio_pct']:.1f}%" if ct['desvio_pct'] else '').font = Font(bold=True)
            next_row += 1

    for col in ws.columns:
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(
            (len(str(c.value or '')) for c in col), default=10
        ) + 4

    resp = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="presupuesto_vs_real_{campaign_name}.xlsx"'
    wb.save(resp)
    return resp


def _build_work_orders_data(request):
    user = request.user
    campaing_id = request.GET.get('campaing_id')
    state_ids = request.GET.getlist('state_ids')
    crop_id = request.GET.get('crop_id')
    area_id = request.GET.get('area_id')
    contractor_id = request.GET.get('contractor_id')
    wo_type = request.GET.get('wo_type')
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')
    producer_filter = _producer_filter(user)

    selected_campaing = None
    orders = []
    all_crops = []
    all_contractors_qs = Contractor.objects.order_by('name')
    if producer_filter is not None:
        all_contractors_qs = all_contractors_qs.filter(producer__in=producer_filter)

    if campaing_id:
        try:
            selected_campaing = Campaing.objects.get(pk=campaing_id)
        except Campaing.DoesNotExist:
            pass

        crops_qs = Crop.objects.filter(campaing_id=campaing_id)
        if producer_filter is not None:
            crops_qs = crops_qs.filter(producer__in=producer_filter)
        all_crops = list(crops_qs.order_by('name'))

        qs = WorkOrder.objects.filter(
            campaing_id=campaing_id,
        ).select_related(
            'crop', 'crop__area', 'stage', 'service', 'stateOT', 'contractor', 'staff', 'area',
        ).order_by('date_init', 'code')

        if producer_filter is not None:
            qs = qs.filter(producer__in=producer_filter)
        if state_ids:
            qs = qs.filter(stateOT_id__in=state_ids)
        if crop_id:
            qs = qs.filter(crop_id=crop_id)
        if area_id:
            qs = qs.filter(area_id=area_id)
        if contractor_id:
            qs = qs.filter(contractor_id=contractor_id)
        if wo_type:
            qs = qs.filter(type=wo_type)
        if date_from:
            qs = qs.filter(date_init__gte=date_from)
        if date_to:
            qs = qs.filter(date_init__lte=date_to)

        orders = list(qs)

    return {
        'campaigns': Campaing.objects.all().order_by('-date_init'),
        'all_states': OTState.objects.all().order_by('name'),
        'all_areas': Area.objects.order_by('name'),
        'all_crops': all_crops,
        'all_contractors': list(all_contractors_qs),
        'wo_types': WorkOrder.types,
        'selected_campaing': selected_campaing,
        'selected_campaing_id': campaing_id,
        'selected_state_ids': state_ids,
        'selected_crop_id': crop_id,
        'selected_area_id': area_id,
        'selected_contractor_id': contractor_id,
        'selected_wo_type': wo_type,
        'selected_date_from': date_from,
        'selected_date_to': date_to,
        'orders': orders,
    }


@login_required
def work_orders_report_view(request):
    ctx = _build_work_orders_data(request)
    return render(request, 'operations/report_work_orders.html', {**adminSite.each_context(request), **ctx})


@login_required
def work_orders_report_excel(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    ctx = _build_work_orders_data(request)
    selected_campaing = ctx['selected_campaing']
    orders = ctx['orders']

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Órdenes de Trabajo'

    header_fill = PatternFill('solid', fgColor='1E7D3B')
    header_font = Font(color='FFFFFF', bold=True, size=9)
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    headers = ['Código', 'Inicio', 'Fin', 'Cultivo', 'Estadio', 'Servicio',
               'Tipo', 'Estado', 'Contratista', 'Costo planeado (MB)', 'Costo real (MB)', 'Personal']

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    TYPE_LABELS = {'TPW': 'Plan técnico', 'OW': 'Propia', 'PW': 'Contratada'}

    for row_idx, wo in enumerate(orders, 2):
        row_data = [
            wo.code,
            wo.date_init.strftime('%d/%m/%Y') if wo.date_init else '',
            wo.date_fin.strftime('%d/%m/%Y') if wo.date_fin else '',
            str(wo.crop) if wo.crop else '',
            str(wo.stage) if wo.stage else '',
            str(wo.service) if wo.service else '',
            TYPE_LABELS.get(wo.type, wo.type) if wo.type else '',
            str(wo.stateOT) if wo.stateOT else '',
            str(wo.contractor) if wo.contractor else '',
            round(float(wo.contractor_cost), 2) if wo.contractor_cost else '',
            round(float(wo.contractor_cost_real), 2) if wo.contractor_cost_real else '',
            str(wo.staff) if wo.staff else '',
        ]
        for col, val in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col, value=val)

    for col in range(1, len(headers) + 1):
        max_len = max(
            (len(str(ws.cell(row=r, column=col).value or '')) for r in range(1, ws.max_row + 1)),
            default=10,
        )
        ws.column_dimensions[get_column_letter(col)].width = min(max_len + 3, 40)

    ws.freeze_panes = 'A2'

    campaign_slug = selected_campaing.name.replace(' ', '_') if selected_campaing else 'sin_campana'
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="ordenes-{campaign_slug}.xlsx"'
    wb.save(response)
    return response
