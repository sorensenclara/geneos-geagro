from user.admin.adminSite import get_producer_for_request
import json
from django.http import HttpResponse, HttpResponseRedirect, JsonResponse
from django.template.loader import get_template
from xhtml2pdf import pisa
from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.db.models import F
from user.models.user_producer import UserProducer
from input.models.purchase_order_detail import PurchaseOrderDetail
from input.models.purchase_order import PurchaseOrder
from input.models.input import Input
from input.models.purchase_reception_detail import PurchaseReceptionDetail
from operation.models.input_work_order import InputWorkOrder
from django.contrib import messages 
from django.utils.translation import gettext_lazy as _
from input.models.type_input import TypeInput
from input.models.input_stock import InputStock
from input.models.stock_movement import StockMovement
from input.models.provider import Provider
from producer.models.store import Store
from django.db.models import Q
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def input_charts_view(request):
    # Muestra gráficos de stock e historial de movimientos para los insumos del usuario.
    # Filtra los insumos según el tipo de usuario (Administración, Organización o Productor).
    # Si se recibe un input_id por GET, genera también un gráfico de evolución acumulada del stock.
    extra_context = {}

    # Filtro de privacidad por productor (mismo criterio que el Admin)
    queryset = Input.objects.all()
    if not request.user.is_superuser and request.user.type != 'Administration':
        if request.user.type == 'Organization':
            user_producers = UserProducer.objects.filter(organization=request.user)
            if user_producers.exists():
                queryset = queryset.filter(producer__in=user_producers)
            else:
                queryset = queryset.none()
        elif request.user.type == 'Producer':
            user_producer = get_producer_for_request(request)
            if user_producer:
                queryset = queryset.filter(producer=user_producer)
            else:
                queryset = queryset.none()
        else:
            queryset = queryset.none()

    all_inputs = queryset.order_by('name')

    # Gráfico General (Barras)
    chart_data_qs = queryset.order_by('-stock')[:15] # Limit to 15 for readability
    extra_context['chart_data_json'] = json.dumps({
        'labels': [i.name for i in chart_data_qs],
        'stock': [i.stock for i in chart_data_qs],
        'projected': [i.stock + i.pending for i in chart_data_qs],
        'min_stock': [getattr(i, 'min_stock', 0) for i in chart_data_qs],
    })

    # Gráfico de Evolución (Solo si se filtra un insumo específico)
    input_id = request.GET.get('input_id')
    extra_context['selected_input_id'] = input_id
    extra_context['selected_input_name'] = _("Select an input") # Default value

    if input_id:
        input_obj = queryset.filter(id=input_id).first()
        if input_obj:
            extra_context['selected_input_name'] = input_obj.name
            
            receptions = PurchaseReceptionDetail.objects.filter(input=input_obj, purchase_reception__state='confirm').values('amount', date=F('purchase_reception__delivery_date'))
            usages = InputWorkOrder.objects.filter(input=input_obj, work_order__stateOT__name="Finalizada").values(amount=F('dose_applied'), date=F('work_order__date_init'))
            
            all_movements = list(receptions)
            for u in usages:
                u['amount'] = float(u['amount'] or 0) * -1
                all_movements.append(u)
            
            movements = sorted([m for m in all_movements if m['date']], key=lambda x: x['date'])
            
            history_labels = [m['date'].strftime('%d/%m/%Y') for m in movements]
            history_values = [sum(float(mov['amount'] or 0) for mov in movements[:i+1]) for i in range(len(movements))] # Calculate cumulative sum
            
            extra_context['timeline_data_json'] = json.dumps({
                'labels': history_labels,
                'data': history_values,
                'name': input_obj.name,
                'min_stock': getattr(input_obj, 'min_stock', 0)
            })
    
    extra_context['all_inputs'] = all_inputs

    from user.admin.adminSite import adminSite
    extra_context.update(adminSite.each_context(request))
    return render(request, 'admin/input/charts.html', extra_context)

def InputPendingPDFView(request):
    pendings = []
    purchase = PurchaseOrder.objects.filter(state="pending")

    if request.user.is_superuser or getattr(request.user, 'type', None) == 'Administration':
        pendings = PurchaseOrderDetail.objects.filter(pending_amount__gt=0, purchase_order__in=purchase)
    elif getattr(request.user, 'type', None) == 'Organization':
        userProducer = UserProducer.objects.filter(organization=request.user)
        if userProducer.exists():
            purchase = purchase.filter(producer__in=userProducer)
            pendings = PurchaseOrderDetail.objects.filter(pending_amount__gt=0, purchase_order__in=purchase)
    elif getattr(request.user, 'type', None) == 'Producer':
        userProducer = get_producer_for_request(request)
        if userProducer:
            purchase = purchase.filter(producer=userProducer)
            pendings = PurchaseOrderDetail.objects.filter(pending_amount__gt=0, purchase_order__in=purchase)
    elif getattr(request.user, 'type', None) == 'Staff':
        userProducer = get_producer_for_request(request)
        if userProducer:
            purchase = purchase.filter(producer=userProducer)
            pendings = PurchaseOrderDetail.objects.filter(pending_amount__gt=0, purchase_order__in=purchase)

    if pendings:
      # Renderiza la plantilla HTML y pasa el queryset como contexto
      context = {'pendientes': pendings}

      template_path = 'purchase/report_purchase_pending.html'
      # Create a Django response object, and specify content_type as pdf
      response = HttpResponse(content_type='application/pdf')
      response['Content-Disposition'] = 'attachment; filename="reporte-compras-pendientes.pdf"'
      # find the template and render it.
      template = get_template(template_path)
      html = template.render(context)

      # create a pdf
      pisa_status = pisa.CreatePDF(html, dest=response)
      # if error then show some funny view
      if pisa_status.err:
         return HttpResponse('Errores: <pre>' + html + '</pre>')
      return response
    else:
        messages.error(request, "No hay pendientes para generar el reporte.")
        return HttpResponseRedirect(request.META.get('HTTP_REFERER', '/'))


def _build_catalog_data(request):

    user = request.user
    type_id = request.GET.get('type_id')
    store_id = request.GET.get('store_id')
    provider_id = request.GET.get('provider_id')
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')
    show_producer = user.is_superuser or getattr(user, 'type', None) == 'Administration'

    # Compute producer filter list (None = no restriction)
    prod_filter = None
    if not user.is_superuser and getattr(user, 'type', None) != 'Administration':
        if getattr(user, 'type', None) == 'Organization':
            prod_filter = list(UserProducer.objects.filter(organization=user))
        elif getattr(user, 'type', None) == 'Producer':
            p = UserProducer.objects.filter(email=user.email).first()
            prod_filter = [p] if p else []
        else:
            prod_filter = []

    def _apply_prod(qs, field='producer'):
        if prod_filter is None:
            return qs
        if not prod_filter:
            return qs.none()
        return qs.filter(**{f'{field}__in': prod_filter})

    # Inputs
    input_qs = Input.objects.select_related('type', 'unit_measure', 'currency', 'producer').order_by('name')
    if prod_filter is None:
        pass
    elif not prod_filter:
        input_qs = input_qs.none()
    else:
        input_qs = input_qs.filter(Q(producer__in=prod_filter) | Q(producer__isnull=True))
    if type_id:
        input_qs = input_qs.filter(type_id=type_id)

    # Stock by store (InputStock)
    stock_qs = InputStock.objects.select_related(
        'input', 'input__type', 'input__unit_measure', 'store'
    ).order_by('input__name', 'store__name')
    stock_qs = _apply_prod(stock_qs, 'store__producer')
    if type_id:
        stock_qs = stock_qs.filter(input__type_id=type_id)
    if store_id:
        stock_qs = stock_qs.filter(store_id=store_id)

    # Movements (StockMovement)
    mov_qs = StockMovement.objects.select_related(
        'input', 'input__type', 'input__unit_measure', 'store',
        'purchase_reception__purchase_order__provider',
        'work_order',
    ).order_by('-date', '-id')
    mov_qs = _apply_prod(mov_qs, 'producer')
    if type_id:
        mov_qs = mov_qs.filter(input__type_id=type_id)
    if store_id:
        mov_qs = mov_qs.filter(store_id=store_id)
    if provider_id:
        mov_qs = mov_qs.filter(
            operation='purchase',
            purchase_reception__purchase_order__provider_id=provider_id,
        )
    if date_from:
        mov_qs = mov_qs.filter(date__gte=date_from)
    if date_to:
        mov_qs = mov_qs.filter(date__lte=date_to)

    raw = list(mov_qs[:501])
    movements_truncated = len(raw) > 500
    movements = raw[:500]

    # Dropdown options
    stores_qs = Store.objects.order_by('name')
    providers_qs = Provider.objects.order_by('name')
    stores_qs = _apply_prod(stores_qs, 'producer')
    providers_qs = _apply_prod(providers_qs, 'producer')

    return {
        'inputs': input_qs,
        'stock_by_store': list(stock_qs),
        'movements': movements,
        'movements_truncated': movements_truncated,
        'types': TypeInput.objects.order_by('name'),
        'all_stores': list(stores_qs),
        'all_providers': list(providers_qs),
        'selected_type_id': type_id,
        'selected_store_id': store_id,
        'selected_provider_id': provider_id,
        'selected_date_from': date_from,
        'selected_date_to': date_to,
        'show_producer': show_producer,
    }


@login_required
def input_catalog_view(request):
    ctx = _build_catalog_data(request)
    from user.admin.adminSite import adminSite
    ctx.update(adminSite.each_context(request))
    return render(request, 'input/report_catalog.html', ctx)


def exchange_rate_api(request):
    # API JSON que devuelve la tasa de conversión de una moneda a la moneda base.
    # Recibe currency_id por GET. Si la moneda ya es la base, devuelve rate=1.0.
    # Usado por el JS del formulario de orden de compra para autocompletar la tasa.
    currency_id = request.GET.get('currency_id')
    if not currency_id:
        return JsonResponse({'error': 'missing currency_id'}, status=400)
    try:
        from parameter.models.currency import Currency
        from parameter.models.exchange_rate import exchangeRate
        currency = Currency.objects.get(pk=currency_id)
        if currency.is_base:
            return JsonResponse({'rate': 1.0, 'is_base': True})
        from agro_system.views.filter_helpers import get_producer_ref
        producer = get_producer_ref(request.user)
        rate = exchangeRate.get_rate(currency, producer=producer)
        return JsonResponse({'rate': rate, 'is_base': False})
    except Currency.DoesNotExist:
        return JsonResponse({'error': 'not found'}, status=404)


def base_currency_api(request):
    from parameter.models.currency import Currency
    currency = Currency.objects.filter(is_base=True).first()
    if currency:
        return JsonResponse({'id': currency.pk, 'name': str(currency)})
    return JsonResponse({'id': None, 'name': ''})


@login_required
def input_catalog_excel_view(request):
    ctx = _build_catalog_data(request)
    show_producer = ctx['show_producer']

    def make_header(ws, headers, fgColor):
        fill = PatternFill('solid', fgColor=fgColor)
        font = Font(color='FFFFFF', bold=True, size=9)
        align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=1, column=col, value=h)
            c.fill = fill; c.font = font; c.alignment = align

    def autowidth(ws, n_cols):
        for col in range(1, n_cols + 1):
            max_len = max(
                (len(str(ws.cell(row=r, column=col).value or '')) for r in range(1, ws.max_row + 1)),
                default=10,
            )
            ws.column_dimensions[get_column_letter(col)].width = min(max_len + 3, 40)

    wb = openpyxl.Workbook()

    # ── Hoja 1: Catálogo ─────────────────────────────────────
    ws1 = wb.active
    ws1.title = 'Catálogo'
    h1 = ['Nombre', 'Tipo', 'Unidad', 'Presentación', 'Precio', 'Moneda',
          'Precio base', 'Modo actualiz.', 'Stock total', 'Stock mín.', 'Pendiente', 'Ctrl. stock']
    if show_producer:
        h1.append('Productor')
    make_header(ws1, h1, '1E7D3B')
    for r, inp in enumerate(ctx['inputs'], 2):
        row = [
            inp.name,
            str(inp.type) if inp.type else '',
            str(inp.unit_measure) if inp.unit_measure else '',
            inp.presentation or '',
            round(float(inp.cost), 2) if inp.cost else 0,
            str(inp.currency) if inp.currency else '',
            round(float(inp.cost_base), 2) if inp.cost_base else '',
            inp.get_cost_update_mode_display() if hasattr(inp, 'get_cost_update_mode_display') else '',
            round(float(inp.stock), 2) if inp.stock else 0,
            round(float(inp.min_stock), 2) if inp.min_stock else 0,
            round(float(inp.pending), 2) if inp.pending else 0,
            'Sí' if inp.control_stock else 'No',
        ]
        if show_producer:
            row.append(str(inp.producer) if inp.producer else '')
        for col, val in enumerate(row, 1):
            ws1.cell(row=r, column=col, value=val)
    ws1.freeze_panes = 'A2'
    autowidth(ws1, len(h1))

    # ── Hoja 2: Stock por almacén ─────────────────────────────
    ws2 = wb.create_sheet('Stock por almacén')
    h2 = ['Insumo', 'Tipo', 'Unidad', 'Almacén', 'Stock']
    make_header(ws2, h2, '1E7D3B')
    for r, s in enumerate(ctx['stock_by_store'], 2):
        ws2.cell(r, 1, str(s.input))
        ws2.cell(r, 2, str(s.input.type) if s.input.type else '')
        ws2.cell(r, 3, str(s.input.unit_measure) if s.input.unit_measure else '')
        ws2.cell(r, 4, str(s.store))
        ws2.cell(r, 5, round(float(s.quantity), 2) if s.quantity else 0)
    ws2.freeze_panes = 'A2'
    autowidth(ws2, len(h2))

    # ── Hoja 3: Movimientos ───────────────────────────────────
    ws3 = wb.create_sheet('Movimientos')
    h3 = ['Fecha', 'Insumo', 'Tipo', 'Unidad', 'Almacén', 'Operación',
          'Cantidad', 'Stock ant.', 'Stock nuevo', 'Proveedor', 'OT', 'Notas']
    make_header(ws3, h3, '2c6e9e')
    for r, m in enumerate(ctx['movements'], 2):
        provider_name = ''
        if m.purchase_reception and m.purchase_reception.purchase_order:
            prov = m.purchase_reception.purchase_order.provider
            provider_name = str(prov) if prov else ''
        ws3.cell(r, 1, str(m.date) if m.date else '')
        ws3.cell(r, 2, str(m.input))
        ws3.cell(r, 3, str(m.input.type) if m.input.type else '')
        ws3.cell(r, 4, str(m.input.unit_measure) if m.input.unit_measure else '')
        ws3.cell(r, 5, str(m.store) if m.store else '')
        ws3.cell(r, 6, m.get_operation_display())
        ws3.cell(r, 7, round(float(m.quantity), 3) if m.quantity else 0)
        ws3.cell(r, 8, round(float(m.stock_before), 3) if m.stock_before is not None else '')
        ws3.cell(r, 9, round(float(m.stock_after), 3) if m.stock_after is not None else '')
        ws3.cell(r, 10, provider_name)
        ws3.cell(r, 11, str(m.work_order) if m.work_order else '')
        ws3.cell(r, 12, m.notes or '')
    ws3.freeze_panes = 'A2'
    autowidth(ws3, len(h3))

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="catalogo-insumos.xlsx"'
    wb.save(response)
    return response


def input_cost_api(request):
    # API JSON que devuelve el costo unitario y la moneda de un insumo dado su ID.
    # Recibe id por GET. Usado por el JS de la orden de trabajo para autocompletar el valor estimado.
    input_id = request.GET.get('id')
    if not input_id:
        return JsonResponse({'error': 'missing id'}, status=400)
    try:
        inp = Input.objects.select_related('currency').get(pk=input_id)
        return JsonResponse({
            'cost': inp.cost,
            'currency': str(inp.currency) if inp.currency else '',
            'currency_id': inp.currency_id,
        })
    except Input.DoesNotExist:
        return JsonResponse({'error': 'not found'}, status=404)