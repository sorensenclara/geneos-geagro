import io
from collections import defaultdict

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import render
from django.db.models import Count, Q, Sum, Avg

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from alert.models.alert import Alert
from alert.models.alert_outcome import AlertOutcome
from alert.models.indicator import Indicator
from operation.models.crop import Crop
from parameter.models.campaign import Campaing
from producer.models.country_side import Countryside
from producer.models.lot import Lot
from producer.models.staff import Staff


# ─── helpers ──────────────────────────────────────────────────────────────────

LEVEL_LABELS = {"warning": "Aviso", "alert": "Alerta", "critical": "Crítico"}
SEVERITY_LABELS = {
    "none": "Sin daño", "minor": "Leve", "moderate": "Moderado",
    "severe": "Grave", "total": "Total",
}
EFFECTIVENESS_LABELS = {
    "yes": "Efectiva", "partial": "Parcial", "no": "No efectiva", "na": "N/A",
}


def _accessible_crops(user):
    if user.type == "Producer":
        return Crop.objects.filter(producer=user)
    if user.type == "Staff":
        staff = Staff.objects.filter(user=user).first()
        return Crop.objects.filter(producer=staff.producer) if staff else Crop.objects.none()
    return Crop.objects.all()


def _build_stats(alerts_qs):
    """Return dict with aggregated stats for a queryset of alerts."""
    total = alerts_qs.count()
    with_outcome = alerts_qs.filter(outcome__isnull=False).count()
    with_damage = alerts_qs.filter(outcome__damage_occurred=True).count()
    effective = alerts_qs.filter(outcome__intervention_effective="yes").count()
    partial = alerts_qs.filter(outcome__intervention_effective="partial").count()
    not_effective = alerts_qs.filter(outcome__intervention_effective="no").count()

    damage_qs = alerts_qs.filter(outcome__damage_occurred=True)
    agg = damage_qs.aggregate(
        area_pct=Avg("outcome__area_affected_pct"),
        loss_kg=Sum("outcome__estimated_loss_kg_ha"),
        loss_ars=Sum("outcome__estimated_loss_ars"),
    )

    return {
        "total": total,
        "with_outcome": with_outcome,
        "without_outcome": total - with_outcome,
        "with_damage": with_damage,
        "effective": effective,
        "partial": partial,
        "not_effective": not_effective,
        "area_pct": round(agg["area_pct"], 1) if agg["area_pct"] is not None else None,
        "loss_kg": round(agg["loss_kg"], 1) if agg["loss_kg"] is not None else None,
        "loss_ars": round(agg["loss_ars"], 0) if agg["loss_ars"] is not None else None,
    }


# ─── view ─────────────────────────────────────────────────────────────────────

@login_required
def CampaignComparisonView(request):
    user = request.user
    crops_base = _accessible_crops(user)
    all_campaigns = Campaing.objects.order_by("-date_init")
    all_indicators = Indicator.objects.order_by("name")

    # ── selected params ────────────────────────────────────────────────────────
    selected_campaign_ids = request.GET.getlist("campaigns")
    group_by = request.GET.get("group_by", "lot")          # "lot" | "variety"
    selected_indicator_ids = request.GET.getlist("indicators")
    download = request.GET.get("download") == "1"

    selected_campaigns = all_campaigns.filter(id__in=selected_campaign_ids) if selected_campaign_ids else Campaing.objects.none()

    # ── base alert queryset ────────────────────────────────────────────────────
    accessible_crop_ids = crops_base.values_list("id", flat=True)
    alerts_base = (
        Alert.objects
        .filter(crop_id__in=accessible_crop_ids, is_aborted=False)
        .select_related("indicator", "crop", "crop__lot", "crop__campaing", "crop__activity")
        .prefetch_related("outcome")
    )

    if selected_campaign_ids:
        alerts_base = alerts_base.filter(crop__campaing_id__in=selected_campaign_ids)
    if selected_indicator_ids:
        alerts_base = alerts_base.filter(indicator_id__in=selected_indicator_ids)

    # ── build comparison matrix ────────────────────────────────────────────────
    # structure: [{group_name, camps: [{camp, stats}]}]
    comparison_rows = []

    if selected_campaign_ids:
        camp_list_q = list(selected_campaigns)
        # collect all group names across all campaigns
        all_groups = set()
        for camp in camp_list_q:
            camp_alerts = alerts_base.filter(crop__campaing=camp)
            if group_by == "variety":
                all_groups.update(camp_alerts.values_list("crop__activity__name", flat=True).distinct())
            else:
                all_groups.update(camp_alerts.values_list("crop__lot__name", flat=True).distinct())

        for group_name in sorted(all_groups):
            camps_data = []
            for camp in camp_list_q:
                camp_alerts = alerts_base.filter(crop__campaing=camp)
                if group_by == "variety":
                    qs = camp_alerts.filter(crop__activity__name=group_name)
                else:
                    qs = camp_alerts.filter(crop__lot__name=group_name)
                camps_data.append({"camp": camp, "stats": _build_stats(qs)})
            comparison_rows.append({"group_name": group_name or ("Sin variedad" if group_by == "variety" else "Sin lote"), "camps": camps_data})

    # ── per-indicator summary per campaign (for the summary table) ─────────────
    indicator_summary = []
    if selected_campaign_ids:
        camp_list_q = list(selected_campaigns)
        for ind in all_indicators:
            camps_data = []
            has_any = False
            for camp in camp_list_q:
                qs = alerts_base.filter(crop__campaing=camp, indicator=ind)
                stats = _build_stats(qs)
                if stats["total"] > 0:
                    has_any = True
                camps_data.append({"camp": camp, "stats": stats})
            if has_any:
                indicator_summary.append({"indicator": ind, "camps": camps_data})

    # ── per-indicator detail: individual alerts with outcomes ──────────────────
    indicator_detail = []
    if selected_campaign_ids:
        for ind in all_indicators:
            ind_alerts = (
                alerts_base.filter(indicator=ind)
                .order_by("crop__campaing__date_init", "date")
                .select_related("crop__lot", "crop__campaing", "crop__activity", "outcome", "outcome__registered_by")
            )
            if ind_alerts.exists():
                indicator_detail.append({"indicator": ind, "alerts": list(ind_alerts)})

    # ── per-crop detail: alerts grouped by lot or variety then by crop ────────
    crop_detail_rows = []
    if selected_campaign_ids:
        all_alerts_detailed = (
            alerts_base
            .order_by("crop__lot__name", "crop__activity__name", "crop__campaing__name", "date")
            .select_related("crop__lot", "crop__campaing", "crop__activity", "indicator", "outcome", "outcome__registered_by")
        )
        groups = {}
        for alert in all_alerts_detailed:
            crop = alert.crop
            group_key = (crop.lot.name if crop.lot else "Sin lote") if group_by != "variety" else (crop.activity.name if crop.activity else "Sin variedad")
            crop_key = crop.pk
            if group_key not in groups:
                groups[group_key] = {"group_name": group_key, "crops": {}}
            if crop_key not in groups[group_key]["crops"]:
                groups[group_key]["crops"][crop_key] = {
                    "variety": crop.activity.name if crop.activity else "—",
                    "lot": crop.lot.name if crop.lot else "—",
                    "campaing": crop.campaing.name if crop.campaing else "—",
                    "alerts": [],
                }
            groups[group_key]["crops"][crop_key]["alerts"].append(alert)
        crop_detail_rows = [
            {"group_name": k, "crops": list(v["crops"].values())}
            for k, v in sorted(groups.items())
        ]

    if download and selected_campaign_ids:
        return _build_excel(selected_campaigns, alerts_base, comparison_rows, indicator_summary, group_by, crop_detail_rows)

    context = {
        "all_campaigns": all_campaigns,
        "all_indicators": all_indicators,
        "selected_campaign_ids": [int(i) for i in selected_campaign_ids],
        "selected_campaigns": selected_campaigns,
        "selected_indicator_ids": [int(i) for i in selected_indicator_ids],
        "group_by": group_by,
        "comparison_rows": comparison_rows,
        "indicator_summary": indicator_summary,
        "indicator_detail": indicator_detail,
        "crop_detail_rows": crop_detail_rows,
        "has_results": bool(selected_campaign_ids),
    }
    return render(request, "pages/report_campaign_comparison.html", context)


# ─── excel builder ─────────────────────────────────────────────────────────────

def _style_header(cell, bg="1F4E79"):
    cell.font = Font(bold=True, color="FFFFFF", size=10)
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _border_all(ws, min_row, max_row, min_col, max_col):
    thin = Side(style="thin", color="CCCCCC")
    b = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = b


_LEVEL_COLORS = {"warning": "FFF9C4", "alert": "FFE0B2", "critical": "FFCDD2"}


def _build_excel(campaigns, alerts_base, comparison_rows, indicator_summary, group_by, crop_detail_rows=None):
    wb = openpyxl.Workbook()

    # ── Sheet 1: Detail by crop (variety + lot) ───────────────────────────────
    ws_crop = wb.active
    group_label_xl = "Variedad" if group_by == "variety" else "Lote"
    sec_label_xl   = "Lote"     if group_by == "variety" else "Variedad"
    ws_crop.title = f"Por {group_label_xl.lower()}"
    crop_headers = [
        group_label_xl, sec_label_xl, "Campaña", "Fecha", "Tipo de alerta", "Nivel",
        "BBCH", "¿Daño?", "Severidad", "Sup. afect. (%)", "Pérd. kg/ha",
        "Pérd. ($)", "Interv. efectiva", "Notas",
    ]
    for c, h in enumerate(crop_headers, start=1):
        _style_header(ws_crop.cell(1, c, h))

    row_num = 2
    for group in (crop_detail_rows or []):
        for crop_row in group["crops"]:
            for alert in crop_row["alerts"]:
                outcome = getattr(alert, "outcome", None)
                ws_crop.cell(row_num, 1, group["group_name"])
                ws_crop.cell(row_num, 2, crop_row["lot"] if group_by == "variety" else crop_row["variety"])
                ws_crop.cell(row_num, 3, crop_row["campaing"])
            ws_crop.cell(row_num, 4, alert.date.strftime("%d/%m/%Y") if alert.date else "")
            ws_crop.cell(row_num, 5, alert.indicator.name if alert.indicator else "")
            ws_crop.cell(row_num, 6, LEVEL_LABELS.get(alert.level, alert.level))
            ws_crop.cell(row_num, 7, alert.bbch_stage or "")
            ws_crop.cell(row_num, 8, "Sí" if (outcome and outcome.damage_occurred) else ("No" if outcome else "Sin reg."))
            ws_crop.cell(row_num, 9, SEVERITY_LABELS.get(outcome.damage_severity, "") if outcome else "")
            ws_crop.cell(row_num, 10, float(outcome.area_affected_pct) if outcome and outcome.area_affected_pct else "")
            ws_crop.cell(row_num, 11, float(outcome.estimated_loss_kg_ha) if outcome and outcome.estimated_loss_kg_ha else "")
            ws_crop.cell(row_num, 12, float(outcome.estimated_loss_ars) if outcome and outcome.estimated_loss_ars else "")
            ws_crop.cell(row_num, 13, EFFECTIVENESS_LABELS.get(outcome.intervention_effective, "") if outcome else "")
            ws_crop.cell(row_num, 14, outcome.notes if outcome else "")
            if outcome and outcome.damage_occurred:
                for c in range(1, 15):
                    ws_crop.cell(row_num, c).fill = PatternFill("solid", fgColor="FFF3E0")
            row_num += 1

    crop_col_widths = [18, 18, 18, 12, 22, 10, 8, 8, 12, 14, 14, 16, 16, 30]
    for i, w in enumerate(crop_col_widths, start=1):
        ws_crop.column_dimensions[get_column_letter(i)].width = w
    ws_crop.auto_filter.ref = f"A1:{get_column_letter(len(crop_headers))}1"
    ws_crop.freeze_panes = "D2"

    # ── Sheet 2: Summary by indicator ─────────────────────────────────────────
    ws1 = wb.create_sheet("Resumen por alerta")
    ws1.row_dimensions[1].height = 30
    ws1.freeze_panes = "C2"

    camp_list = list(campaigns)
    col_offset = 3  # col A=Indicador, B=Criticidad, then one group per campaign
    COLS_PER_CAMP = 7

    # headers
    ws1.cell(1, 1, "Tipo de alerta").font = Font(bold=True, size=11)
    ws1.cell(1, 2, "Criticidad").font = Font(bold=True, size=11)
    _style_header(ws1.cell(1, 1))
    _style_header(ws1.cell(1, 2))

    sub_labels = ["Total", "Con daño", "Sup. afect. (%)", "Pérd. kg/ha", "Pérd. ($)", "Interv. efectiva", "Sin registro"]
    for i, camp in enumerate(camp_list):
        col = col_offset + i * COLS_PER_CAMP
        ws1.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + COLS_PER_CAMP - 1)
        hdr = ws1.cell(1, col, camp.name)
        _style_header(hdr, bg="2E75B6")
        for j, label in enumerate(sub_labels):
            _style_header(ws1.cell(2, col + j, label), bg="9DC3E6")

    ws1.row_dimensions[2].height = 20

    crit_colors = {"low": "E8F5E9", "medium": "FFF9C4", "high": "FFCDD2"}
    for r, row in enumerate(indicator_summary, start=3):
        ind = row["indicator"]
        ws1.cell(r, 1, ind.name)
        crit_cell = ws1.cell(r, 2, ind.get_criticality_level_display())
        crit_cell.fill = PatternFill("solid", fgColor=crit_colors.get(ind.criticality_level, "FFFFFF"))
        for i, camp_data in enumerate(row["camps"]):
            col = col_offset + i * COLS_PER_CAMP
            stats = camp_data["stats"]
            ws1.cell(r, col,     stats.get("total", 0))
            ws1.cell(r, col + 1, stats.get("with_damage", 0))
            ws1.cell(r, col + 2, stats.get("area_pct") or "")
            ws1.cell(r, col + 3, stats.get("loss_kg") or "")
            ws1.cell(r, col + 4, stats.get("loss_ars") or "")
            ws1.cell(r, col + 5, stats.get("effective", 0) + stats.get("partial", 0))
            ws1.cell(r, col + 6, stats.get("without_outcome", 0))
            if stats.get("with_damage", 0) > 0:
                ws1.cell(r, col + 1).fill = PatternFill("solid", fgColor="FFCDD2")
            if stats.get("loss_ars"):
                ws1.cell(r, col + 4).fill = PatternFill("solid", fgColor="FFCDD2")

    total_cols = col_offset + len(camp_list) * COLS_PER_CAMP - 1
    _border_all(ws1, 1, max(len(indicator_summary) + 2, 3), 1, total_cols)
    ws1.column_dimensions["A"].width = 28
    ws1.column_dimensions["B"].width = 12
    for i in range(len(camp_list) * COLS_PER_CAMP):
        ws1.column_dimensions[get_column_letter(col_offset + i)].width = 14

    # ── Sheet 3: Raw alert detail ──────────────────────────────────────────────
    ws3 = wb.create_sheet("Detalle de alertas")
    headers = [
        "Campaña", "Fecha", "Tipo de alerta", "Lote", "Variedad",
        "Nivel", "BBCH", "¿Daño?", "Severidad", "Área afectada (%)",
        "Pérdida estimada (kg/ha)", "Pérdida económica ($)", "Efectividad", "Notas",
    ]
    for c, h in enumerate(headers, start=1):
        _style_header(ws3.cell(1, c, h))

    for r, alert in enumerate(alerts_base.order_by("crop__campaing__date_init", "date"), start=2):
        outcome = getattr(alert, "outcome", None)
        ws3.cell(r, 1, alert.crop.campaing.name if alert.crop else "")
        ws3.cell(r, 2, alert.date.strftime("%d/%m/%Y") if alert.date else "")
        ws3.cell(r, 3, alert.indicator.name if alert.indicator else "")
        ws3.cell(r, 4, alert.crop.lot.name if alert.crop and alert.crop.lot else "")
        ws3.cell(r, 5, alert.crop.activity.name if alert.crop and alert.crop.activity else "")
        ws3.cell(r, 6, LEVEL_LABELS.get(alert.level, alert.level))
        ws3.cell(r, 7, alert.bbch_stage or "")
        ws3.cell(r, 8, "Sí" if (outcome and outcome.damage_occurred) else "No")
        ws3.cell(r, 9, SEVERITY_LABELS.get(outcome.damage_severity, "") if outcome else "")
        ws3.cell(r, 10, float(outcome.area_affected_pct) if outcome and outcome.area_affected_pct else "")
        ws3.cell(r, 11, float(outcome.estimated_loss_kg_ha) if outcome and outcome.estimated_loss_kg_ha else "")
        ws3.cell(r, 12, float(outcome.estimated_loss_ars) if outcome and outcome.estimated_loss_ars else "")
        ws3.cell(r, 13, EFFECTIVENESS_LABELS.get(outcome.intervention_effective, "") if outcome else "Sin registro")
        ws3.cell(r, 14, outcome.notes if outcome else "")

        if outcome and outcome.damage_occurred:
            for c in range(1, 15):
                ws3.cell(r, c).fill = PatternFill("solid", fgColor="FFF3E0")

    col_widths = [18, 12, 22, 18, 18, 10, 8, 8, 12, 14, 20, 20, 14, 30]
    for i, w in enumerate(col_widths, start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    ws3.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    # ── stream response ────────────────────────────────────────────────────────
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    campaign_names = "_vs_".join(c.name.replace(" ", "-") for c in camp_list)
    filename = f"informe_alertas_{campaign_names}.xlsx"
    response = HttpResponse(
        buf.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
